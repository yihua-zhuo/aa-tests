#!/usr/bin/env python3
"""Local OCR sidecar for attachment extraction.

Runs one engine per process so PaddleOCR and Docling can live in separate
virtualenvs:

    ~/.venvs/pulse-paddle/bin/python scripts/ocr_sidecar.py \
        --engine paddle --port 8765 --upload-root data/uploads
    ~/.venvs/pulse-docling/bin/python scripts/ocr_sidecar.py \
        --engine docling --port 8766 --upload-root data/uploads

The server only accepts absolute paths under --upload-root, binds to loopback
by default, and optionally requires the X-OCR-Token header (OCR_TOKEN).

Extraction is native-first: digital PDFs use pdfplumber text blocks and
tables, Office files use python-docx/openpyxl/python-pptx (Docling remains the
fallback), and only scanned pages are rendered for OCR. Every response also
carries a lightweight Document IR so the worker can assess quality and route
recovery attempts without re-parsing the file.
"""

from __future__ import annotations

import argparse
import difflib
import hashlib
import hmac
import html
import json
import math
import os
import platform
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, NamedTuple

os.environ.setdefault("PADDLE_PDX_MODEL_SOURCE", "bos")
os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")

PADDLE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".pdf",
}
DOCLING_OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
# Docling can also read page images and PDFs. The standard pipeline enables OCR
# by default (`do_ocr=True`, auto OCR engine) and the sidecar configures it
# explicitly, so this engine can serve as an independent whole-page candidate.
# GIF is excluded because Docling's image backends do not accept it.
DOCLING_EXTENSIONS = DOCLING_OFFICE_EXTENSIONS | (PADDLE_EXTENSIONS - {".gif"})
# Bump when extraction post-processing changes so cached text is not reused.
PIPELINE_VERSION = "dir2"
MAX_BODY_BYTES = 64 * 1024
MAX_TEXT_CHARS = 1_000_000
MAX_DIR_CHARS = 1_000_000
MAX_DIR_BLOCKS = 2_000
MAX_BLOCK_CHARS = 20_000
TEXT_LAYER_MIN_CHARS = 20
TEXT_LAYER_TRUST_MIN = 0.6
TEXT_CROSSCHECK_RATE = 0.1
TEXT_CROSSCHECK_MAX_PAGES = 2
TEXT_CROSSCHECK_MIN_SIMILARITY = 0.85
DEFAULT_PDF_DPI = 200
DEFAULT_MAX_PDF_PAGES = 100
DEFAULT_MAX_PDF_TOTAL_PAGES = 1000
DEFAULT_PDF_TIME_BUDGET_SECONDS = 240
MAX_IMAGE_FRAMES = 20
MAX_IMAGE_PIXELS = 80_000_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_OFFICE_ENTRIES = 10_000
MAX_OFFICE_RATIO = 500
XML_PROLOGUE_BYTES = 4096
DEFAULT_MAX_OFFICE_PAGES = 200
DEFAULT_CAMELOT_MAX_PAGES = 200
DEFAULT_CAMELOT_TIME_BUDGET_SECONDS = 120
DEFAULT_DOCLING_MAX_PAGES = 500
# Ink-gap miss detection: binarize the page, mask every box the recognizer
# already claimed, then keep the ink no box explains.
DEFAULT_INK_MAX_SIDE = 900
DEFAULT_INK_MIN_PIXELS = 45
DEFAULT_INK_MAX_GAPS = 12
# More than this share of the page ink left uncovered means the boxes do not
# line up with the pixels (rotated page, broken coordinates), not that the
# recognizer missed something.
INK_UNSAFE_UNCOVERED_RATIO = 0.6
# A cell reading one of these still counts as "no text": its ruled borders are
# masked, but ink sitting inside it is a miss. Mirrors PLACEHOLDER_CELL in
# src/ocr-quality.ts.
INK_BLANK_CELL_TEXT = re.compile(
    r"^(?:[-—–/\\]|n/?a|无|暂无|待定|未知|none|null)$",
    re.IGNORECASE,
)
OCR_MODES = {"auto", "table", "structure", "enhanced"}
OCR_CHANNELS = {"rgb", "gray", "red"}
OCR_LANGS = {"ch", "en", "japan", "korean", "arabic", "ru"}

# CJK ideographs plus the kana/hangul blocks Paddle can emit for mixed text.
CJK_CLASS = r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\u3040-\u30ff\uac00-\ud7af"
LATIN_OR_DIGIT = r"A-Za-z0-9"


class RequestError(Exception):
    def __init__(self, status: int, message: str) -> None:
        super().__init__(message)
        self.status = status
        self.message = message


class Extraction(NamedTuple):
    text: str
    confidence: float | None
    # Display label; PDFs with a mixed text layer gain a `+pdf-text` suffix.
    model: str
    # Base engine model used to validate cached results across restarts.
    base_model: str
    pages: int
    truncated: bool
    route: str = "visual"
    document: dict[str, Any] | None = None


class PageExtraction(NamedTuple):
    text: str
    confidence: float | None
    page: dict[str, Any]
    model: str
    preprocess: dict[str, Any] | None = None


class NativeText(NamedTuple):
    text: str
    pages: list[dict[str, Any]]
    truncated: bool = False
    meta: dict[str, Any] | None = None


def env_int(name: str, default: int) -> int:
    try:
        return max(1, int(os.environ.get(name, str(default))))
    except ValueError:
        return default


def env_int_zero_ok(name: str, default: int) -> int:
    """Knob where 0 means "turn this off"; `env_int`'s floor of 1 would swallow it."""
    try:
        return max(0, int(os.environ.get(name, str(default))))
    except ValueError:
        return default


def image_budget_issue(
    width: int,
    height: int,
    frames: int,
    max_pixels: int = MAX_IMAGE_PIXELS,
    max_frames: int = MAX_IMAGE_FRAMES,
) -> tuple[int, str] | None:
    """Status and message for an image that exceeds the decode budget, else None."""
    if width <= 0 or height <= 0:
        return 422, "image has invalid dimensions"
    # Frames are decoded one at a time, so the budget applies per frame.
    if width * height > max_pixels:
        return 413, f"image frame exceeds the {max_pixels // 1_000_000} megapixel limit"
    if max(1, frames) > max_frames:
        return 413, f"image has more than {max_frames} frames"
    return None


def validate_image_budget(path: Path) -> None:
    """Reject decompression bombs before Pillow/numpy allocate the full bitmap."""
    try:
        from PIL import Image
    except Exception:  # noqa: BLE001 - Pillow is optional
        return
    # Guards every later decode, including engines that open the file themselves.
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(path) as image:
            width, height = image.size
            frames = int(getattr(image, "n_frames", 1) or 1)
    except RequestError:
        raise
    except Exception as cause:  # noqa: BLE001 - let the engine report malformed inputs
        raise RequestError(422, f"cannot read image: {cause}") from None
    issue = image_budget_issue(
        int(width), int(height), frames, MAX_IMAGE_PIXELS, env_int("OCR_MAX_IMAGE_FRAMES", MAX_IMAGE_FRAMES)
    )
    if issue is not None:
        raise RequestError(issue[0], issue[1])


def validate_office_archive(path: Path) -> None:
    """Bound OOXML decompression so a small upload cannot exhaust memory."""
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_OFFICE_ENTRIES:
                raise RequestError(413, "Office archive contains too many parts")
            total = 0
            for info in infos:
                total += info.file_size
                if total > MAX_OFFICE_UNCOMPRESSED_BYTES:
                    raise RequestError(413, "Office archive expands beyond the size limit")
                if info.file_size > 5_000_000:
                    ratio = info.file_size / max(1, info.compress_size)
                    if ratio > MAX_OFFICE_RATIO:
                        raise RequestError(413, "Office archive has a suspicious compression ratio")
                if info.filename.lower().endswith((".xml", ".rels")):
                    # Only the declaration can nest entities, so the header is enough.
                    header = archive.open(info).read(XML_PROLOGUE_BYTES)
                    reject_explosive_xml(info.filename, header)
    except RequestError:
        raise
    except zipfile.BadZipFile:
        raise RequestError(422, "invalid Office archive") from None
    except OSError as cause:
        raise RequestError(422, f"cannot read Office archive: {cause}") from None


def reject_explosive_xml(name: str, header: bytes) -> None:
    """Refuse OOXML parts that declare entities: an entity bomb expands in the parser."""
    if b"<!DOCTYPE" in header or b"<!ENTITY" in header:
        raise RequestError(422, f"Office part {name} declares XML entities")


def pdf_needs_password(path: Path) -> bool:
    """True when a PDF cannot be opened without a non-empty password.

    Encrypted documents otherwise surface as "no text recognized", which the
    queue retries three times and an operator reads as a model failure.
    """
    try:
        from pdfminer.pdfdocument import PDFDocument
        from pdfminer.pdfparser import PDFParser
    except Exception:  # noqa: BLE001 - pdfminer ships with pdfplumber
        return False
    try:
        with path.open("rb") as handle:
            PDFDocument(PDFParser(handle), password="")
    except Exception as cause:  # noqa: BLE001 - only the two password errors matter
        return type(cause).__name__ in {"NoPasswordError", "PDFPasswordIncorrect"}
    return False


def env_float(name: str, default: float) -> float:
    try:
        return float(os.environ.get(name, str(default)))
    except ValueError:
        return default


_RUNTIME_LABEL: str | None = None


def runtime_label() -> str:
    """Paddle/PaddleOCR versions feed the cache key so upgrades invalidate reuse."""
    global _RUNTIME_LABEL
    if _RUNTIME_LABEL is not None:
        return _RUNTIME_LABEL
    parts: list[str] = []
    for module_name in ("paddleocr", "paddle", "torch"):
        try:
            module = __import__(module_name)
            version = getattr(module, "__version__", None)
            if version:
                parts.append(f"{module_name}-{version}")
        except Exception:  # noqa: BLE001 - optional runtime
            continue
    _RUNTIME_LABEL = "|".join(parts)
    return _RUNTIME_LABEL


def runtime_details() -> dict[str, str | None]:
    """Split runtime versions for result metadata and model-fingerprint keys."""
    paddle_version: str | None = None
    cuda_version: str | None = None
    try:
        import paddle  # type: ignore

        paddle_version = getattr(paddle, "__version__", None)
        if getattr(paddle, "device", None) is not None:
            try:
                if paddle.device.is_compiled_with_cuda():
                    cuda_version = str(paddle.version.cuda() or "cuda")
            except Exception:  # noqa: BLE001 - CPU builds may lack version.cuda
                cuda_version = None
    except Exception:  # noqa: BLE001 - optional runtime
        pass
    return {
        "python": platform.python_version(),
        "paddle": paddle_version,
        "cuda": cuda_version,
    }


def parse_model_profiles() -> dict[str, dict[str, str | None]]:
    """Named detector/recognizer pairs for model comparison and canary runs.

    ``OCR_MODEL_PROFILES`` holds JSON such as
    ``{"paddle-v6-server": {"det": "PP-OCRv6_server_det", "rec": "PP-OCRv6_server_rec"}}``.
    Profiles are opt-in: an unknown name is rejected instead of silently
    falling back to the default model.
    """
    raw = os.environ.get("OCR_MODEL_PROFILES", "").strip()
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except ValueError:
        print("OCR_MODEL_PROFILES is not valid JSON; ignoring", file=sys.stderr)
        return {}
    if not isinstance(parsed, dict):
        return {}
    profiles: dict[str, dict[str, str | None]] = {}
    for key, value in parsed.items():
        if not isinstance(key, str) or not key.strip() or not isinstance(value, dict):
            continue
        det = value.get("det")
        rec = value.get("rec")
        profiles[key.strip()] = {
            "det": det.strip() if isinstance(det, str) and det.strip() else None,
            "rec": rec.strip() if isinstance(rec, str) and rec.strip() else None,
        }
    return profiles


_MODEL_SHA256_CACHE: dict[tuple[str, ...], str | None] = {}


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


_SCRIPT_PATTERNS = (
    ("japan", re.compile(r"[\u3040-\u30ff\u31f0-\u31ff]")),
    ("korean", re.compile(r"[\u1100-\u11ff\u3130-\u318f\uac00-\ud7af]")),
    ("arabic", re.compile(r"[\u0600-\u06ff\u0750-\u077f\ufb50-\ufdff\ufe70-\ufeff]")),
    ("ru", re.compile(r"[\u0400-\u04ff\u0500-\u052f]")),
    ("ch", re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")),
)


def detect_language(text: str) -> str | None:
    """Pick an OCR language from already-extracted native text.

    Kana/Hangul/Arabic/Cyrillic pages need their own recognizer; Han text keeps
    the Chinese model even when Latin abbreviations are mixed in. Returns
    ``None`` for pure Latin text so the default model stays in place.
    """
    sample = text[:50_000]
    counts = {name: len(pattern.findall(sample)) for name, pattern in _SCRIPT_PATTERNS}
    for name in ("japan", "korean", "arabic", "ru"):
        if counts[name] >= 3:
            return name
    if counts["ch"] >= 3:
        return "ch"
    return None


def model_files_sha256(model_names: list[str]) -> str | None:
    """SHA256 over the model files actually loaded.

    ``OCR_MODEL_SHA256`` can pin the value. Otherwise ``OCR_MODEL_DIR``
    (path-separated) and the default PaddleX cache are scanned for directories
    matching the detector/recognizer names. Returns ``None`` when nothing can be
    found, so callers never mistake a name-only key for a verified hash.
    """
    names = tuple(sorted(name for name in model_names if name))
    if names in _MODEL_SHA256_CACHE:
        return _MODEL_SHA256_CACHE[names]
    override = os.environ.get("OCR_MODEL_SHA256", "").strip()
    if override:
        _MODEL_SHA256_CACHE[names] = override
        return override
    roots = [
        Path(item)
        for item in os.environ.get("OCR_MODEL_DIR", "").split(os.pathsep)
        if item.strip()
    ]
    roots.extend(
        [
            Path.home() / ".paddlex" / "official_models",
            Path.home() / ".paddleocr",
            Path.home() / ".cache" / "paddle" / "official_models",
        ]
    )
    wanted = [name.lower() for name in names]
    candidates: list[Path] = []
    for root in roots:
        if not root.is_dir():
            continue
        for model_dir in sorted(root.iterdir()):
            if not model_dir.is_dir():
                continue
            lowered = model_dir.name.lower()
            if wanted and not any(name in lowered or lowered in name for name in wanted):
                continue
            for file in sorted(model_dir.rglob("*")):
                if file.is_file() and file.suffix.lower() in {".pdmodel", ".pdiparams", ".json", ".yml", ".yaml", ".txt"}:
                    candidates.append(file)
    if not candidates:
        _MODEL_SHA256_CACHE[names] = None
        return None
    digest = hashlib.sha256()
    for file in sorted(candidates)[:200]:
        try:
            digest.update(str(file.relative_to(file.parents[1])).encode("utf-8"))
            digest.update(_hash_file(file).encode("ascii"))
        except OSError:
            continue
    _MODEL_SHA256_CACHE[names] = digest.hexdigest()
    return _MODEL_SHA256_CACHE[names]


def estimate_skew(gray: Any) -> float | None:
    """Median near-horizontal line angle; ``None`` without OpenCV/numpy."""
    try:
        import cv2
        import numpy as np
    except Exception:  # noqa: BLE001 - optional dependency
        return None
    try:
        sample = gray.copy()
        sample.thumbnail((800, 800))
        array = np.asarray(sample, dtype=np.uint8)
        edges = cv2.Canny(array, 50, 150, apertureSize=3)
        lines = cv2.HoughLinesP(
            edges,
            1,
            np.pi / 180,
            threshold=80,
            minLineLength=max(40, sample.width // 4),
            maxLineGap=8,
        )
        if lines is None:
            return None
        angles: list[float] = []
        for x1, y1, x2, y2 in lines[:, 0]:
            angle = math.degrees(math.atan2(float(y2 - y1), float(x2 - x1)))
            if -45 <= angle <= 45:
                angles.append(angle)
        if not angles:
            return None
        angles.sort()
        return round(angles[len(angles) // 2], 2)
    except Exception:  # noqa: BLE001 - skew is advisory only
        return None


def _colorfulness(rgb: Any) -> float:
    try:
        import numpy as np
    except Exception:  # noqa: BLE001 - optional dependency
        return 0.0
    try:
        data = np.asarray(rgb.resize((128, 128)), dtype=np.float32)
        red, green, blue = data[..., 0], data[..., 1], data[..., 2]
        rg = red - green
        yb = 0.5 * (red + green) - blue
        value = math.sqrt(float(rg.std()) ** 2 + float(yb.std()) ** 2) + 0.3 * math.sqrt(
            float(rg.mean()) ** 2 + float(yb.mean()) ** 2
        )
        return min(1.0, value / 110.0)
    except Exception:  # noqa: BLE001 - optional dependency
        return 0.0


def image_metrics(image: Any) -> dict[str, Any] | None:
    """Blur, exposure, shadow, glare, noise, color and orientation signals."""
    try:
        from PIL import Image, ImageChops, ImageFilter, ImageStat
    except Exception:  # noqa: BLE001 - Pillow is optional
        return None
    try:
        width, height = image.size
        if width == 0 or height == 0:
            return None
        rgb = image.convert("RGB")
        gray = rgb.convert("L")
        sample = gray.copy()
        sample.thumbnail((1000, 1000))
        edges = sample.filter(ImageFilter.FIND_EDGES)
        # Edge variance correlates with sharpness; normalize to 0-1 so the
        # TypeScript quality thresholds are stable across resolutions.
        sharpness = min(1.0, float(ImageStat.Stat(edges).var[0]) / 5000.0)
        stats = ImageStat.Stat(sample)
        brightness = float(stats.mean[0]) / 255.0
        contrast = float(stats.stddev[0]) / 255.0

        # Shadow = variation of per-cell paper brightness. Taking a high
        # percentile inside each cell ignores text strokes, so a clean page
        # scores near zero while a fold/shadow gradient lights up.
        shadow = 0.0
        try:
            import numpy as np

            array = np.asarray(sample, dtype=np.float32)
            rows, columns = 8, 8
            y_edges = np.linspace(0, array.shape[0], rows + 1).astype(int)
            x_edges = np.linspace(0, array.shape[1], columns + 1).astype(int)
            paper: list[float] = []
            for row in range(rows):
                for column in range(columns):
                    cell = array[
                        y_edges[row] : y_edges[row + 1],
                        x_edges[column] : x_edges[column + 1],
                    ]
                    if cell.size:
                        paper.append(float(np.percentile(cell, 90)))
            if len(paper) >= 4:
                shadow = min(
                    1.0,
                    max(
                        0.0,
                        (float(np.percentile(paper, 95)) - float(np.percentile(paper, 5)))
                        / 255.0,
                    ),
                )
        except Exception:  # noqa: BLE001 - numpy optional
            small = sample.resize((16, 16), Image.BOX)
            values = sorted(float(value) for value in small.getdata())
            low = values[int(len(values) * 0.1)]
            high = values[int(len(values) * 0.9)]
            shadow = min(1.0, max(0.0, (high - low) / 255.0))

        # Specular glare: near-white pixels clearly brighter than the page's
        # dominant paper tone. Pure-white paper (mode >= 245) is not glare,
        # otherwise every clean document would be flagged.
        histogram = sample.histogram()
        total = max(1, sum(histogram))
        background = max(range(256), key=histogram.__getitem__)
        if background >= 245:
            glare = 0.0
        else:
            cutoff = max(250, background + 25)
            glare = sum(histogram[cutoff:]) / total

        # Noise = residual after median smoothing; text edges survive the median.
        residual = ImageChops.difference(sample, sample.filter(ImageFilter.MedianFilter(3)))
        noise = min(1.0, float(ImageStat.Stat(residual).stddev[0]) / 255.0 * 2.0)

        skew = estimate_skew(gray)
        try:
            exif = int(image.getexif().get(274, 1) or 1)
        except Exception:  # noqa: BLE001 - malformed EXIF
            exif = None
        rotation = {3: 180, 6: 90, 8: 270}.get(exif or 1, 0)

        issues: list[str] = []
        if sharpness < 0.15:
            issues.append("blurry")
        if brightness < 0.25:
            issues.append("underexposed")
        elif brightness > 0.95:
            issues.append("overexposed")
        if contrast < 0.08:
            issues.append("low_contrast")
        if width * height < 300_000:
            issues.append("low_resolution")
        if glare >= 0.02:
            issues.append("glare")
        if shadow >= 0.35:
            issues.append("shadow")
        if skew is not None and abs(skew) >= 3:
            issues.append("skewed")
        if noise >= 0.09:
            issues.append("noisy")

        retake_reasons: list[str] = []
        if glare >= 0.08:
            retake_reasons.append("glare")
        if brightness < 0.12 or brightness > 0.98:
            retake_reasons.append("exposure")
        if sharpness < 0.03 and contrast < 0.08:
            retake_reasons.append("blurry")
        if shadow >= 0.6:
            retake_reasons.append("shadow")
        if width * height < 120_000:
            retake_reasons.append("low_resolution")

        return {
            "width": width,
            "height": height,
            "sharpness": round(sharpness, 4),
            "brightness": round(brightness, 4),
            "contrast": round(contrast, 4),
            "shadow": round(shadow, 4),
            "glare": round(glare, 4),
            "noise": round(noise, 4),
            "colorfulness": round(_colorfulness(rgb), 4),
            "skew": skew,
            "orientation": {
                "exif": exif,
                "applied": rotation,
                "aspect": round(width / height, 4) if height else None,
                "uncertain": skew is None,
            },
            "issues": issues,
            "retake": {"required": bool(retake_reasons), "reasons": retake_reasons},
        }
    except Exception:  # noqa: BLE001 - preflight must never block OCR
        return None


def preflight_image(path: Path) -> dict[str, Any] | None:
    """Cheap quality signals; never fails extraction."""
    try:
        from PIL import Image
    except Exception:  # noqa: BLE001 - Pillow is optional
        return None
    try:
        with Image.open(path) as image:
            metrics = image_metrics(image)
            if metrics is None:
                return None
            dpi_value: float | None = None
            dpi_info = (getattr(image, "info", {}) or {}).get("dpi")
            if isinstance(dpi_info, (tuple, list)) and dpi_info:
                try:
                    dpi_value = round(float(dpi_info[0]), 1) or None
                except (TypeError, ValueError):
                    dpi_value = None
        metrics["dpi"] = dpi_value
        return metrics
    except Exception:  # noqa: BLE001 - preflight must never block OCR
        return None


def normalize_ocr_text(text: str) -> str:
    """Space and width folding for display, never for field values.

    Rules that could move a character inside an identifier, a clock, a URL or a
    list are disabled by masking those spans first; the un-normalized engine
    output is still returned alongside as `rawText`.
    """
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    # Docling's markdown export escapes "<" as "&lt;" and OCR can echo entities
    # from source documents; plain-text extraction should return the glyph.
    text = html.unescape(text).replace("\u00a0", " ")
    masked, tokens = _mask_spans(text, PROTECTED_SPANS)
    # Paddle's Chinese recognizer often returns full-width separators inside
    # numbers (1，280，500) — only in strict thousands grouping, so a list like
    # "1，2，3" keeps its punctuation.
    masked = re.sub(
        r"(?<![\d，])(\d{1,3})((?:，\d{3})+)",
        lambda match: match.group(1) + match.group(2).replace("，", ","),
        masked,
    )
    masked = re.sub(r"(?<=\d)\s*．\s*(?=\d)", ".", masked)
    masked = re.sub(r"(?<=\d)\s*％", "%", masked)
    masked = re.sub(r"(?<=\d)\s*－\s*(?=\d)", "-", masked)
    # Codes such as X1-2 or AB123 are field values: never split them with a space.
    masked, code_tokens = _mask_code_tokens(masked)
    tokens.update(code_tokens)
    # Mixed-script lines often lose the separator ("新增客户42家",
    # "客户Customer"). Restore a single space at CJK/Latin boundaries.
    masked = re.sub(rf"(?<=[{CJK_CLASS}])(?=[{LATIN_OR_DIGIT}%])", " ", masked)
    masked = re.sub(rf"(?<=[{LATIN_OR_DIGIT}%])(?=[{CJK_CLASS}])", " ", masked)
    # "Customer:上海" reads as "Customer: 上海"; sentence punctuation before
    # Latin text needs the same separator.
    masked = re.sub(rf"(?<=[:：])(?=[{CJK_CLASS}])", " ", masked)
    masked = re.sub(rf"(?<=[。！？；])(?=[{LATIN_OR_DIGIT}])", " ", masked)
    return _restore_spans(masked, tokens)


def _mask_spans(text: str, patterns: tuple[re.Pattern[str], ...]) -> tuple[str, dict[str, str]]:
    """Replace protected spans with private-use placeholders and return the map."""
    tokens: dict[str, str] = {}
    for pattern in patterns:
        def hold(match: re.Match[str]) -> str:
            key = f"\ue000{len(tokens):04d}\ue001"
            tokens[key] = match.group(0)
            return key

        text = pattern.sub(hold, text)
    return text, tokens


def _mask_code_tokens(text: str) -> tuple[str, dict[str, str]]:
    """Hold alphanumeric runs that mix letters and digits (identifier-shaped)."""
    tokens: dict[str, str] = {}

    def hold(match: re.Match[str]) -> str:
        run = match.group(0)
        has_letter = any(character.isascii() and character.isalpha() for character in run)
        has_digit = any(character.isdigit() for character in run)
        if not (has_letter and has_digit):
            return run
        key = f"\ue000{len(tokens):04d}\ue001"
        tokens[key] = run
        return key

    return re.sub(rf"[{LATIN_OR_DIGIT}]+", hold, text), tokens


def _restore_spans(text: str, tokens: dict[str, str]) -> str:
    for key, value in tokens.items():
        text = text.replace(key, value)
    return text


PROTECTED_SPANS: tuple[re.Pattern[str], ...] = (
    re.compile(r"(?:https?|ftp|mailto)://\S+|www\.\S+", re.IGNORECASE),
    re.compile(rf"\d{{1,2}}[:：]\d{{2}}(?:[:：]\d{{2}})?"),
    re.compile(rf"[:：][ \u00a0]*[{LATIN_OR_DIGIT}][{LATIN_OR_DIGIT}\-_.·/#]+"),
)


def text_layer_trust(text: str) -> tuple[float, list[str]]:
    """Score a PDF text layer in [0, 1] and list the reasons for any penalty.

    Native text is the cheapest path but not always trustworthy: scanned pages
    can carry a near-empty placeholder layer, broken embedded fonts emit
    replacement/private-use glyphs, and OCR'd archives embed garbage. Score the
    layer before preferring it over a rendered OCR pass.
    """
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return 0.0, ["empty"]
    reasons: list[str] = []
    length_score = min(1.0, len(compact) / TEXT_LAYER_MIN_CHARS)
    if len(compact) < TEXT_LAYER_MIN_CHARS:
        reasons.append("short")
    replacement = text.count("\ufffd")
    control = sum(1 for char in text if ord(char) < 32 and char not in "\n\r\t")
    private = sum(1 for char in text if 0xE000 <= ord(char) <= 0xF8FF)
    # Measure against non-whitespace so a long, sparse page cannot dilute garbage.
    garbage_ratio = (replacement + control + private) / len(compact)
    garbage_score = max(0.0, 1.0 - garbage_ratio * 10)
    if garbage_ratio > 0:
        reasons.append("garbage")
    readable_ratio = sum(1 for char in compact if char.isalnum()) / len(compact)
    readable_score = min(1.0, readable_ratio / 0.7)
    if readable_ratio < 0.7:
        reasons.append("symbols")
    return round(max(0.0, min(length_score, garbage_score, readable_score)), 3), reasons


def text_layer_threshold() -> float:
    return min(1.0, max(0.0, env_float("OCR_PDF_TEXT_TRUST_MIN", TEXT_LAYER_TRUST_MIN)))


def usable_text_layer(text: str) -> bool:
    score, _ = text_layer_trust(text)
    return score >= text_layer_threshold()


def page_text_trust(page: dict[str, Any]) -> tuple[float, list[str]]:
    stored = page.get("trust")
    if isinstance(stored, (int, float)):
        return float(stored), list(page.get("trustReasons") or [])
    return text_layer_trust(str(page.get("text") or ""))


def trustworthy_native_page(page: dict[str, Any]) -> bool:
    score, _ = page_text_trust(page)
    return score >= text_layer_threshold()


def text_layer_summary(page_trusts: list[tuple[float, list[str]]]) -> dict[str, Any]:
    if not page_trusts:
        return {
            "pages": 0,
            "trusted": 0,
            "threshold": text_layer_threshold(),
            "minTrust": None,
            "meanTrust": None,
            "reasons": {},
        }
    threshold = text_layer_threshold()
    reasons: dict[str, int] = {}
    for _, page_reasons in page_trusts:
        for reason in page_reasons:
            reasons[reason] = reasons.get(reason, 0) + 1
    scores = [score for score, _ in page_trusts]
    return {
        "pages": len(scores),
        "trusted": sum(1 for score in scores if score >= threshold),
        "threshold": threshold,
        "minTrust": round(min(scores), 3),
        "meanTrust": round(sum(scores) / len(scores), 3),
        "reasons": reasons,
    }


def text_crosscheck_rate() -> float:
    return min(1.0, max(0.0, env_float("OCR_PDF_TEXT_CROSSCHECK_RATE", TEXT_CROSSCHECK_RATE)))


def text_crosscheck_max_pages() -> int:
    return env_int("OCR_PDF_TEXT_CROSSCHECK_MAX_PAGES", TEXT_CROSSCHECK_MAX_PAGES)


def text_crosscheck_min_similarity() -> float:
    return min(
        1.0,
        max(
            0.0,
            env_float("OCR_PDF_TEXT_CROSSCHECK_MIN_SIMILARITY", TEXT_CROSSCHECK_MIN_SIMILARITY),
        ),
    )


def select_crosscheck_pages(
    trusted: list[int],
    rate: float,
    max_pages: int,
    priority: list[int] | None = None,
) -> list[int]:
    """Deterministic sample that always includes risk-heavy trusted pages.

    Amount/number pages are read even when the global rate would skip them;
    the rest of the sample stays evenly spaced and stable across runs.
    """
    if not trusted or rate <= 0 or max_pages <= 0:
        return []
    wanted = min(max_pages, max(1, math.ceil(len(trusted) * rate)))
    if wanted >= len(trusted):
        return list(trusted)
    priority_order = [page for page in (priority or []) if page in trusted]
    risk_quota = min(len(priority_order), max(0, math.ceil(max_pages / 2)))
    selected = list(priority_order[:risk_quota])
    remaining = wanted - len(selected)
    if remaining <= 0:
        return sorted(selected)[:max_pages]
    if remaining == 1:
        candidates = [page for page in trusted if page not in selected]
        if candidates:
            selected.append(candidates[len(candidates) // 2])
    else:
        step = (len(trusted) - 1) / (remaining - 1)
        for index in range(remaining):
            selected.append(trusted[round(index * step)])
    return sorted(set(selected))[:max_pages]


CROSSCHECK_COMPARE_CHARS = 20_000
NUMERIC_TOKEN = re.compile(r"\d+(?:[.,，．]\d+)*")
THOUSANDS_GROUP = re.compile(r"^\d{1,3}(?:,\d{3})+(?:\.\d+)?$")


def crosscheck_numeric_tokens(text: str) -> list[str]:
    normalized = text.translate(
        str.maketrans("０１２３４５６７８９，．", "0123456789,.")
    )
    tokens = []
    for token in NUMERIC_TOKEN.findall(normalized):
        value = token.replace(",", "") if THOUSANDS_GROUP.match(token) else token
        if value:
            tokens.append(value)
    return sorted(tokens)


CORRUPT_MARKERS = re.compile(r"[?？□�]")


def photometric_selection(
    enhanced: tuple[str, float | None],
    baseline: tuple[str, float | None],
) -> dict[str, Any]:
    """Pick an enhanced reading only when a non-score signal supports it.

    The same recognizer produces both candidates, so a higher confidence alone
    is self-confirmation. Business tokens that the original preserved and the
    candidate dropped are treated as a conflict and keep the original text.
    """
    enhanced_text, enhanced_confidence = enhanced
    baseline_text, baseline_confidence = baseline
    enhanced_tokens = Counter(crosscheck_numeric_tokens(enhanced_text))
    baseline_tokens = Counter(crosscheck_numeric_tokens(baseline_text))
    baseline_total = sum(baseline_tokens.values())
    shared = sum((enhanced_tokens & baseline_tokens).values())
    missing = sum((baseline_tokens - enhanced_tokens).values())
    added = sum((enhanced_tokens - baseline_tokens).values())
    agreement = 1.0 if baseline_total == 0 else round(shared / baseline_total, 3)
    enhanced_corrupt = len(CORRUPT_MARKERS.findall(enhanced_text))
    baseline_corrupt = len(CORRUPT_MARKERS.findall(baseline_text))
    enhanced_chars = len(enhanced_text.strip())
    baseline_chars = len(baseline_text.strip())
    delta = (enhanced_confidence or 0.0) - (baseline_confidence or 0.0)
    numeric_conflict = baseline_total > 0 and agreement < 0.8 and baseline_corrupt == 0

    if numeric_conflict:
        selected, basis = "original", "numeric_conflict"
    elif enhanced_corrupt < baseline_corrupt and enhanced_chars >= baseline_chars * 0.98:
        selected, basis = "enhanced", "corrupt"
    elif added > 0 and agreement >= 0.9 and enhanced_chars >= baseline_chars * 0.98:
        selected, basis = "enhanced", "tokens"
    elif delta >= 0.02 and agreement >= 0.95 and enhanced_chars >= baseline_chars * 0.98:
        selected, basis = "enhanced", "confidence"
    else:
        selected, basis = "original", "no_gain"

    return {
        "selected": selected,
        "basis": basis,
        "conflict": numeric_conflict,
        "agreement": agreement,
        "enhanced": {
            "confidence": enhanced_confidence,
            "characters": enhanced_chars,
            "numericTokens": sum(enhanced_tokens.values()),
            "corrupt": enhanced_corrupt,
            "missingTokens": missing,
            "addedTokens": added,
        },
        "original": {
            "confidence": baseline_confidence,
            "characters": baseline_chars,
            "numericTokens": baseline_total,
            "corrupt": baseline_corrupt,
        },
    }


def crosscheck_priority_pages(native_pages: list[dict[str, Any]], positions: list[int]) -> list[int]:
    """Trusted pages with the most numeric content, highest density first."""
    ranked = []
    for position in positions:
        if position >= len(native_pages):
            continue
        tokens = crosscheck_numeric_tokens(str(native_pages[position].get("text") or ""))
        if tokens:
            ranked.append((len(tokens), position))
    ranked.sort(key=lambda item: (-item[0], item[1]))
    return [position for _, position in ranked]


def text_similarity(expected: str, observed: str) -> float:
    """Bounded, whitespace-insensitive similarity in [0, 1] for the crosscheck sample."""
    left = re.sub(r"\s+", "", expected)[:CROSSCHECK_COMPARE_CHARS]
    right = re.sub(r"\s+", "", observed)[:CROSSCHECK_COMPARE_CHARS]
    if not left and not right:
        return 1.0
    if not left or not right:
        return 0.0
    return round(difflib.SequenceMatcher(None, left, right, autojunk=False).ratio(), 3)


def crosscheck_record(
    number: int, expected: str, observed: str, threshold: float
) -> dict[str, Any]:
    similarity = text_similarity(expected, observed)
    expected_numbers = crosscheck_numeric_tokens(expected)
    numeric_mismatch = bool(
        expected_numbers and expected_numbers != crosscheck_numeric_tokens(observed)
    )
    return {
        "page": number,
        "similarity": similarity,
        "matched": similarity >= threshold,
        "numericMismatch": numeric_mismatch,
    }


def crosscheck_summary(
    samples: list[dict[str, Any]], threshold: float, planned: int, failed: int = 0
) -> dict[str, Any]:
    similarities = [item["similarity"] for item in samples]
    return {
        "sampled": len(samples),
        "planned": planned,
        "failed": failed,
        "matched": sum(1 for item in samples if item["matched"]),
        "mismatched": sum(1 for item in samples if not item["matched"]),
        "numericMismatched": sum(1 for item in samples if item.get("numericMismatch")),
        "threshold": threshold,
        "minSimilarity": round(min(similarities), 3) if similarities else None,
        "pages": samples,
        "budgetExceeded": planned > len(samples) + failed,
    }


def model_label(det: str, rec: str) -> str:
    combined = f"{det} {rec}".lower()
    if "v6_medium" in combined:
        return "PP-OCRv6-medium"
    if "v6_small" in combined:
        return "PP-OCRv6-small"
    if "v6_tiny" in combined:
        return "PP-OCRv6-tiny"
    if "v5_server" in combined:
        return "PP-OCRv5-server"
    if det == "default":
        return "PP-OCRv5-default"
    return f"{det}/{rec}"


def is_cjk(char: str) -> bool:
    return "\u3400" <= char <= "\u9fff" or "\uf900" <= char <= "\ufaff"


def join_line_fragments(items: list[tuple[str, float, float, float, float]]) -> str:
    line = items[0][0]
    for previous, current in zip(items, items[1:]):
        both_cjk = is_cjk(previous[0][-1:]) and is_cjk(current[0][:1])
        line += ("" if both_cjk else " ") + current[0]
    return line


def box_bbox(x1: float, y1: float, x2: float, y2: float) -> dict[str, float]:
    return {"x": x1, "y": y1, "width": max(x2 - x1, 0.0), "height": max(y2 - y1, 0.0)}


def bbox_from_value(value: Any) -> dict[str, float] | None:
    if isinstance(value, (list, tuple)) and len(value) >= 4:
        try:
            return box_bbox(float(value[0]), float(value[1]), float(value[2]), float(value[3]))
        except (TypeError, ValueError):
            return None
    if isinstance(value, dict):
        x = value.get("x")
        y = value.get("y")
        width = value.get("width")
        height = value.get("height")
        if all(isinstance(item, (int, float)) for item in (x, y, width, height)):
            return box_bbox(float(x), float(y), float(x) + float(width), float(y) + float(height))
    return None


def dir_source(engine: str, model: str | None, stage: str) -> dict[str, Any]:
    return {"engine": engine, "model": model or None, "stage": stage}


def dir_block(
    block_id: str,
    block_type: str,
    text: str,
    bbox: dict[str, float] | None,
    confidence: float | None,
    source: dict[str, Any],
    order: int,
    rows: list[list[str]] | None = None,
    cells: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    block: dict[str, Any] = {
        "id": block_id,
        "type": block_type,
        "text": text[:MAX_BLOCK_CHARS],
        "bbox": bbox,
        "confidence": confidence,
        "source": source,
        "validation": None,
        "order": order,
    }
    if rows:
        block["rows"] = [
            [str(cell)[:2000] for cell in row[:100]] for row in rows[:1000]
        ]
    if cells:
        block["cells"] = cells[:1000]
    return block


def dir_page(
    number: int,
    width: float | None,
    height: float | None,
    blocks: list[dict[str, Any]],
    source: dict[str, Any] | None,
    meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    page = {
        "number": number,
        "width": width,
        "height": height,
        "blocks": blocks,
        "source": source,
    }
    if meta:
        page["meta"] = meta
    return page


def median_text_height(pages: list[dict[str, Any]]) -> float | None:
    heights: list[float] = []
    for page in pages:
        for block in page.get("blocks") or []:
            if block.get("type") not in {"paragraph", "heading"}:
                continue
            bbox = block.get("bbox") or {}
            height = bbox.get("height") if isinstance(bbox, dict) else None
            if isinstance(height, (int, float)) and height > 0:
                heights.append(float(height))
    if not heights:
        return None
    heights.sort()
    return round(heights[len(heights) // 2], 2)


def dir_document(
    route: str, pages: list[dict[str, Any]], meta: dict[str, Any] | None = None
) -> dict[str, Any]:
    kept_pages: list[dict[str, Any]] = []
    reading_order: list[str] = []
    total_chars = 0
    total_blocks = 0
    dropped = False
    for page in pages:
        kept_blocks: list[dict[str, Any]] = []
        for block in page["blocks"]:
            if total_blocks >= MAX_DIR_BLOCKS or total_chars >= MAX_DIR_CHARS:
                dropped = True
                break
            chars = len(block.get("text") or "")
            if total_chars + chars > MAX_DIR_CHARS:
                dropped = True
                continue
            kept_blocks.append(block)
            total_chars += chars
            total_blocks += 1
            if (block.get("text") or "").strip():
                reading_order.append(block["id"])
        kept_pages.append({**page, "blocks": kept_blocks})
    document_meta = dict(meta or {})
    if dropped:
        document_meta["truncated"] = True
    if "textHeight" not in document_meta:
        text_height = median_text_height(kept_pages)
        if text_height is not None:
            document_meta["textHeight"] = text_height
    return {
        "version": 1,
        "route": route,
        "pages": kept_pages,
        "readingOrder": reading_order,
        "meta": document_meta,
    }


def raw_text_from_document(document: Any) -> str | None:
    """Engine text before display normalization, rebuilt from the DIR blocks.

    Only the aggregate `text` is normalized, so the un-normalized copy is
    recovered from block payloads instead of threading a second string through
    every pipeline. Layout separators may differ from `text`.
    """
    if not isinstance(document, dict):
        return None
    pages = document.get("pages")
    if not isinstance(pages, list):
        return None
    page_texts: list[str] = []
    for page in pages:
        blocks = page.get("blocks") if isinstance(page, dict) else None
        if not isinstance(blocks, list):
            continue
        block_texts = [
            str(block.get("text") or "")
            for block in blocks
            if isinstance(block, dict) and str(block.get("text") or "").strip()
        ]
        if block_texts:
            page_texts.append("\n".join(block_texts))
    return "\n\n".join(page_texts) if page_texts else None


ScoredEntry = tuple[str, float | None, float, float, float, float]


def group_scored_boxes(
    texts: list[str],
    scores: list[float],
    boxes: Any,
    max_gap_ratio: float | None = 1.5,
) -> list[list[ScoredEntry]]:
    entries: list[ScoredEntry] = []
    for index, (text, box) in enumerate(zip(texts, boxes)):
        if not text.strip():
            continue
        x1, y1, x2, y2 = (float(part) for part in list(box)[:4])
        score = scores[index] if index < len(scores) else None
        entries.append((text, score, x1, y1, x2, y2))
    entries.sort(key=lambda item: ((item[3] + item[5]) / 2, item[2]))

    groups: list[list[ScoredEntry]] = []
    bounds: list[tuple[float, float]] = []
    for entry in entries:
        height = max(entry[5] - entry[3], 1.0)
        for index in range(len(groups) - 1, -1, -1):
            top, bottom = bounds[index]
            group_height = max(bottom - top, 1.0)
            overlap = min(bottom, entry[5]) - max(top, entry[3])
            right = max(item[4] for item in groups[index])
            gap = entry[2] - right
            fits = overlap >= 0.5 * min(height, group_height)
            if max_gap_ratio is not None:
                fits = fits and gap <= max_gap_ratio * max(height, group_height)
            if fits:
                groups[index].append(entry)
                bounds[index] = (min(top, entry[3]), max(bottom, entry[5]))
                break
        else:
            groups.append([entry])
            bounds.append((entry[3], entry[5]))
    return [sorted(group, key=lambda item: item[2]) for group in groups]


def scored_cells(groups: list[list[ScoredEntry]]) -> list[dict[str, Any]]:
    cells: list[dict[str, Any]] = []
    for row_index, group in enumerate(groups):
        for column_index, entry in enumerate(group):
            cell: dict[str, Any] = {
                "row": row_index,
                "column": column_index,
                "text": entry[0][:2000],
                "bbox": box_bbox(entry[2], entry[3], entry[4], entry[5]),
            }
            if entry[1] is not None:
                cell["confidence"] = round(float(entry[1]), 4)
            cells.append(cell)
    return cells


def normalized_cell_text(value: Any) -> str:
    return "".join(str(value).split())


def table_pred_cells(
    grid: list[list[str]], table: Any
) -> list[dict[str, Any]] | None:
    """Per-cell scores from PP-StructureV3 when its OCR order lines up with the grid."""
    pred = table.get("table_ocr_pred") if hasattr(table, "get") else None
    if not hasattr(pred, "get"):
        return None
    texts = [str(text) for text in (pred.get("rec_texts") or [])]
    if not texts:
        return None
    positions = [
        (row_index, column_index)
        for row_index, row in enumerate(grid)
        for column_index, cell in enumerate(row)
        if str(cell).strip()
    ]
    if len(positions) != len(texts):
        return None
    matching = sum(
        1
        for index, (row_index, column_index) in enumerate(positions)
        if normalized_cell_text(grid[row_index][column_index])
        == normalized_cell_text(texts[index])
    )
    if matching * 2 < len(texts):
        return None
    scores = [float(score) for score in (pred.get("rec_scores") or [])]
    boxes = (
        pred.get("cell_box_list")
        or pred.get("rec_boxes")
        or pred.get("rec_polys")
    )
    cells: list[dict[str, Any]] = []
    for index, (row_index, column_index) in enumerate(positions):
        cell: dict[str, Any] = {
            "row": row_index,
            "column": column_index,
            "text": texts[index][:2000],
        }
        if index < len(scores):
            cell["confidence"] = round(scores[index], 4)
        if boxes is not None and index < len(boxes):
            bbox = bbox_from_value(boxes[index])
            if bbox:
                cell["bbox"] = bbox
        cells.append(cell)
    return cells or None


FORMULA_GLYPHS = set("=+-−×÷±∑∫√∞≈≠≤≥^_/()[]{}<>|%²³")
HEADING_END = set("。！？；.!?;:：,，")
LAYOUT_LABELS = {
    "doc_title": ("heading", "title"),
    "paragraph_title": ("heading", "heading"),
    "title": ("heading", "heading"),
    "text": ("paragraph", "text"),
    "list": ("paragraph", "list"),
    "aside_text": ("paragraph", "aside"),
    "footnote": ("paragraph", "footnote"),
    "number": ("paragraph", "number"),
    "table": ("table", "table"),
    "figure": ("chart", "figure"),
    "chart": ("chart", "chart"),
    "formula": ("formula", "formula"),
    "seal": ("stamp", "stamp"),
    "stamp": ("stamp", "stamp"),
    "signature": ("signature", "signature"),
    "handwriting": ("handwriting", "handwriting"),
    "header": ("header", "header"),
    "footer": ("footer", "footer"),
    "image": ("image", "image"),
}


def classify_text_blocks(
    blocks: list[dict[str, Any]], page_height: float | None
) -> None:
    """Heuristic heading/formula/footnote hints for line-grouped OCR blocks."""
    heights = [
        float(block["bbox"]["height"])
        for block in blocks
        if isinstance(block.get("bbox"), dict) and block["bbox"].get("height")
    ]
    heights.sort()
    median = heights[len(heights) // 2] if heights else None
    for block in blocks:
        if block.get("type") != "paragraph":
            continue
        text = str(block.get("text") or "").strip()
        if not text:
            continue
        glyphs = sum(1 for char in text if char in FORMULA_GLYPHS)
        formula_ratio = glyphs / max(1, len(text))
        bbox = block.get("bbox") if isinstance(block.get("bbox"), dict) else None
        region: str | None = None
        if formula_ratio >= 0.25 and glyphs >= 2 and len(text) <= 160:
            block["type"] = "formula"
            region = "formula"
        elif (
            median
            and bbox
            and float(bbox.get("height") or 0) >= median * 1.5
            and len(text) <= 80
            and text[-1] not in HEADING_END
        ):
            block["type"] = "heading"
            region = "heading"
        elif (
            median
            and bbox
            and page_height
            and float(bbox.get("y") or 0) >= page_height * 0.88
            and len(text) <= 160
        ):
            region = "footnote"
        if region:
            block["meta"] = {"region": region, "regionSource": "heuristic"}


def paddle_page_blocks(
    result: Any,
    page_number: int,
    stage: str,
    engine_label: str,
    model_label_value: str,
    table_hint: bool = False,
    table_bbox: dict[str, float] | None = None,
) -> tuple[str, float | None, list[dict[str, Any]], float | None, float | None]:
    texts = [str(text) for text in (result.get("rec_texts") or [])]
    scores = [float(score) for score in (result.get("rec_scores") or [])]
    boxes = result.get("rec_boxes")
    source = dir_source(engine_label, model_label_value, stage)
    overall = sum(scores) / len(scores) if scores else None

    if boxes is None or len(boxes) != len(texts):
        text = "\n".join(texts)
        blocks = (
            [dir_block(f"p{page_number}-b0", "paragraph", text, None, overall, source, 0)]
            if text.strip()
            else []
        )
        return text, overall, blocks, None, None

    if table_hint and table_bbox is not None:
        # Region-level routing: only entries inside the ruled table area go
        # through table grouping. Titles, notes and prose keep line grouping.
        table_texts: list[str] = []
        table_scores: list[float] = []
        table_boxes: list[Any] = []
        body_texts: list[str] = []
        body_scores: list[float] = []
        body_boxes: list[Any] = []
        for text, score, box in zip(texts, scores, boxes):
            raw_box = list(box)[:4]
            entry_bbox = bbox_from_value(raw_box)
            if entry_bbox is not None and _center_inside_bbox(entry_bbox, table_bbox):
                table_texts.append(text)
                table_scores.append(score)
                table_boxes.append(raw_box)
            else:
                body_texts.append(text)
                body_scores.append(score)
                body_boxes.append(raw_box)
        if table_texts:
            groups = group_scored_boxes(
                table_texts, table_scores, table_boxes, max_gap_ratio=None
            )
            rows = [[cell[0] for cell in group] for group in groups]
            table_text = "\n".join(" | ".join(row) for row in rows)
            cell_boxes = [(cell[2], cell[3], cell[4], cell[5]) for group in groups for cell in group]
            block_bbox = box_bbox(
                min(cell[0] for cell in cell_boxes),
                min(cell[1] for cell in cell_boxes),
                max(cell[2] for cell in cell_boxes),
                max(cell[3] for cell in cell_boxes),
            )
            blocks = [
                dir_block(
                    f"p{page_number}-t0",
                    "table",
                    table_text,
                    block_bbox,
                    overall,
                    source,
                    0,
                    rows=rows,
                    cells=scored_cells(groups),
                )
            ]
            body_groups = group_scored_boxes(body_texts, body_scores, body_boxes)
            for index, group in enumerate(body_groups, 1):
                fragments = [(item[0], item[2], item[3], item[4], item[5]) for item in group]
                group_scores = [item[1] for item in group if item[1] is not None]
                blocks.append(
                    dir_block(
                        f"p{page_number}-b{index}",
                        "paragraph",
                        join_line_fragments(fragments),
                        box_bbox(
                            min(item[2] for item in group),
                            min(item[3] for item in group),
                            max(item[4] for item in group),
                            max(item[5] for item in group),
                        ),
                        sum(group_scores) / len(group_scores) if group_scores else None,
                        source,
                        index,
                    )
                )
            width = max(
                (block.get("bbox") or {}).get("x", 0.0)
                + (block.get("bbox") or {}).get("width", 0.0)
                for block in blocks
            )
            height = max(
                (block.get("bbox") or {}).get("y", 0.0)
                + (block.get("bbox") or {}).get("height", 0.0)
                for block in blocks
            )
            blocks.sort(
                key=lambda block: (
                    (block.get("bbox") or {}).get("y", float("inf")),
                    (block.get("bbox") or {}).get("x", 0.0),
                )
            )
            for index, block in enumerate(blocks):
                block["order"] = index
            classify_text_blocks(blocks, height or None)
            text = "\n".join(block["text"] for block in blocks)
            return text, overall, blocks, width or None, height or None

    if table_hint:
        groups = group_scored_boxes(texts, scores, boxes, max_gap_ratio=None)
        rows = [[cell[0] for cell in group] for group in groups]
        width = max((max(cell[4] for cell in group) for group in groups), default=0.0)
        height = max((max(cell[5] for cell in group) for group in groups), default=0.0)
        text = "\n".join(" | ".join(row) for row in rows)
        bbox = box_bbox(0.0, 0.0, width, height) if rows else None
        blocks = [
            dir_block(
                f"p{page_number}-t0",
                "table",
                text,
                bbox,
                overall,
                source,
                0,
                rows=rows,
                cells=scored_cells(groups),
            )
        ] if text.strip() else []
        return text, overall, blocks, width or None, height or None

    groups = group_scored_boxes(texts, scores, boxes)
    blocks = []
    width = 0.0
    height = 0.0
    for index, group in enumerate(groups):
        fragments = [(item[0], item[2], item[3], item[4], item[5]) for item in group]
        block_text = join_line_fragments(fragments)
        group_scores = [item[1] for item in group if item[1] is not None]
        bbox = box_bbox(
            min(item[2] for item in group),
            min(item[3] for item in group),
            max(item[4] for item in group),
            max(item[5] for item in group),
        )
        width = max(width, bbox["x"] + bbox["width"])
        height = max(height, bbox["y"] + bbox["height"])
        blocks.append(
            dir_block(
                f"p{page_number}-b{index}",
                "paragraph",
                block_text,
                bbox,
                sum(group_scores) / len(group_scores) if group_scores else None,
                source,
                index,
            )
        )
    classify_text_blocks(blocks, height or None)
    text = "\n".join(block["text"] for block in blocks)
    return text, overall, blocks, width or None, height or None


def detect_table_grid(path: Path) -> dict[str, float] | None:
    """Bbox of the ruled table area in image pixels, when one is found."""
    try:
        import cv2
    except Exception:  # noqa: BLE001 - OpenCV ships with PaddleOCR but is optional
        return None
    try:
        image = cv2.imread(str(path), cv2.IMREAD_GRAYSCALE)
        if image is None:
            return None
        height, width = image.shape[:2]
        if width < 240 or height < 120:
            return None
        scale = min(1.0, 1200.0 / max(width, height))
        if scale < 1.0:
            image = cv2.resize(
                image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA
            )
            height, width = image.shape[:2]
        binary = cv2.adaptiveThreshold(
            image,
            255,
            cv2.ADAPTIVE_THRESH_MEAN_C,
            cv2.THRESH_BINARY_INV,
            15,
            10,
        )
        horizontal = cv2.morphologyEx(
            binary,
            cv2.MORPH_OPEN,
            cv2.getStructuringElement(cv2.MORPH_RECT, (max(width // 25, 10), 1)),
        )
        vertical = cv2.morphologyEx(
            binary,
            cv2.MORPH_OPEN,
            cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(height // 8, 10))),
        )
        horizontal_boxes = cv2.findContours(
            horizontal, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )[0]
        vertical_boxes = cv2.findContours(
            vertical, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )[0]
        long_h = [
            cv2.boundingRect(box)
            for box in horizontal_boxes
            if cv2.boundingRect(box)[2] > width * 0.5
        ]
        long_v = [
            cv2.boundingRect(box)
            for box in vertical_boxes
            if cv2.boundingRect(box)[3] > height * 0.4
        ]
        if len(long_h) < 2 or len(long_v) < 2:
            return None
        rects = long_h + long_v
        left = min(rect[0] for rect in rects)
        top = min(rect[1] for rect in rects)
        right = max(rect[0] + rect[2] for rect in rects)
        bottom = max(rect[1] + rect[3] for rect in rects)
        return box_bbox(left / scale, top / scale, right / scale, bottom / scale)
    except Exception:  # noqa: BLE001 - detection must never fail an extraction
        return None

def table_html_grid(pred_html: str) -> list[list[str]]:
    grid: list[list[str]] = []
    for row in re.findall(r"<tr[^>]*>(.*?)</tr>", pred_html, flags=re.S | re.I):
        cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, flags=re.S | re.I)
        cleaned = [
            html.unescape(re.sub(r"<[^>]+>", "", cell)).strip() for cell in cells
        ]
        if any(cleaned):
            grid.append(cleaned)
    return grid




def temp_image_path() -> Path:
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
        return Path(handle.name)


def crop_region_image(
    path: Path,
    extension: str,
    page_number: int,
    bbox: dict[str, Any],
    dpi: int | None,
    padding: float,
    bbox_dpi: int | None = None,
) -> Path:
    """Crop one bbox (pixel coordinates of the bbox's render) with context.

    The crop is upscaled when short so small amounts, dates and identifiers get
    enough pixels for the recognizer. Raises RequestError for invalid input.
    """
    try:
        from PIL import Image
    except Exception:  # noqa: BLE001 - region recovery requires Pillow
        raise RequestError(501, "Pillow is required for region extraction") from None

    cleanup: list[Path] = []
    document = None
    bbox_scale = 1.0
    try:
        if extension == ".pdf":
            try:
                import pypdfium2 as pdfium
            except Exception:  # noqa: BLE001 - optional PDF renderer
                raise RequestError(501, "pypdfium2 is required for PDF region extraction") from None
            document = pdfium.PdfDocument(str(path))
            total = len(document)
            if page_number < 1 or page_number > total:
                raise RequestError(400, f"page must be between 1 and {total}")
            page = document[page_number - 1]
            dpi_value = dpi or env_int("OCR_PDF_DPI", DEFAULT_PDF_DPI)
            if bbox_dpi and bbox_dpi != dpi_value:
                bbox_scale = dpi_value / bbox_dpi
            try:
                page_width, page_height = page.get_size()
                pixels = (float(page_width) * dpi_value / 72) * (
                    float(page_height) * dpi_value / 72
                )
                if pixels > MAX_IMAGE_PIXELS:
                    raise RequestError(
                        413,
                        f"rendered page exceeds the {MAX_IMAGE_PIXELS // 1_000_000} megapixel limit",
                    )
            except RequestError:
                raise
            except Exception:  # noqa: BLE001 - size metadata is optional
                pass
            bitmap = page.render(scale=dpi_value / 72)
            try:
                image = bitmap.to_pil().convert("RGB")
            finally:
                bitmap.close()
        else:
            frames, cleanup, _ = prepare_image_frames(path)
            index = min(max(page_number, 1), len(frames)) - 1
            with Image.open(frames[index]) as opened:
                image = opened.convert("RGB")

        width, height = image.size
        try:
            x = float(bbox["x"])
            y = float(bbox["y"])
            box_width = float(bbox["width"])
            box_height = float(bbox["height"])
        except (KeyError, TypeError, ValueError):
            raise RequestError(400, "bbox must contain numeric x, y, width and height") from None
        if not all(value == value and abs(value) != float("inf") for value in (x, y, box_width, box_height)):
            raise RequestError(400, "bbox values must be finite")
        if box_width <= 0 or box_height <= 0:
            raise RequestError(400, "bbox width and height must be positive")
        if bbox_scale != 1.0:
            x *= bbox_scale
            y *= bbox_scale
            box_width *= bbox_scale
            box_height *= bbox_scale

        margin = max(0.0, min(padding, 1.0)) * box_height
        left = max(0, int(x - margin))
        top = max(0, int(y - margin))
        right = min(width, int(x + box_width + margin + 0.5))
        bottom = min(height, int(y + box_height + margin + 0.5))
        if right - left < 2 or bottom - top < 2:
            raise RequestError(400, "bbox is outside the page or too small")

        crop = image.crop((left, top, right, bottom))
        if crop.height < 64:
            scale = 64.0 / max(crop.height, 1)
            crop = crop.resize(
                (max(1, round(crop.width * scale)), 64), Image.LANCZOS
            )
        temporary = temp_image_path()
        crop.save(temporary, format="PNG")
        return temporary
    finally:
        for temp in cleanup:
            temp.unlink(missing_ok=True)
        if document is not None:
            try:
                document.close()
            except Exception:  # noqa: BLE001 - already closed
                pass


def prepare_image_frames(path: Path) -> tuple[list[Path], list[Path], bool]:
    """Apply EXIF orientation and split multi-frame TIFF/GIF into PNG frames."""
    cleanup: list[Path] = []
    try:
        from PIL import Image, ImageOps
    except Exception:  # noqa: BLE001 - Pillow is optional; Paddle handles plain images
        return [path], cleanup, False
    try:
        with Image.open(path) as image:
            frames = int(getattr(image, "n_frames", 1) or 1)
            if frames <= 1:
                orientation = image.getexif().get(274, 1)
                if orientation not in (None, 1):
                    temp = temp_image_path()
                    ImageOps.exif_transpose(image).convert("RGB").save(temp, format="PNG")
                    cleanup.append(temp)
                    return [temp], cleanup, False
                return [path], cleanup, False
            limit = min(frames, env_int("OCR_MAX_IMAGE_FRAMES", MAX_IMAGE_FRAMES))
            out: list[Path] = []
            for index in range(limit):
                image.seek(index)
                temp = temp_image_path()
                image.convert("RGB").save(temp, format="PNG")
                out.append(temp)
                cleanup.append(temp)
            return out, cleanup, frames > limit
    except Exception:  # noqa: BLE001 - malformed metadata should not fail extraction
        for temp in cleanup:
            temp.unlink(missing_ok=True)
        return [path], [], False


def _otsu_threshold(gray: Any) -> int:
    histogram = gray.histogram()
    total = sum(histogram)
    if total == 0:
        return 128
    sum_total = sum(index * count for index, count in enumerate(histogram))
    weight_background = 0
    sum_background = 0
    best_variance = 0.0
    threshold = 128
    for index, count in enumerate(histogram):
        weight_background += count
        if weight_background == 0:
            continue
        weight_foreground = total - weight_background
        if weight_foreground == 0:
            break
        sum_background += index * count
        mean_background = sum_background / weight_background
        mean_foreground = (sum_total - sum_background) / weight_foreground
        variance = (
            weight_background
            * weight_foreground
            * (mean_background - mean_foreground) ** 2
        )
        if variance > best_variance:
            best_variance = variance
            threshold = index
    return threshold


def _shadow_corrected(image: Any) -> Any:
    """Divide by a low-frequency background estimate to flatten shadows."""
    from PIL import Image, ImageChops, ImageFilter, ImageOps

    gray = image.convert("L")
    radius = max(8, min(gray.size) // 20)
    background = gray.filter(ImageFilter.GaussianBlur(radius))
    try:
        import numpy as np

        data = np.asarray(gray, dtype=np.float32)
        base = np.asarray(background, dtype=np.float32)
        normalized = np.clip(data / np.maximum(base, 1.0) * 200.0, 0, 255).astype(np.uint8)
        return Image.fromarray(normalized, mode="L")
    except Exception:  # noqa: BLE001 - PIL fallback
        flattened = ImageChops.subtract(gray, background, scale=1, offset=255)
        return ImageOps.autocontrast(flattened, cutoff=1)


def _contrast_enhanced(image: Any) -> Any:
    from PIL import Image, ImageFilter, ImageOps

    try:
        import cv2
        import numpy as np

        gray = np.asarray(image.convert("L"))
        enhanced = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
        base = Image.fromarray(enhanced)
    except Exception:  # noqa: BLE001 - PIL fallback
        base = ImageOps.autocontrast(image.convert("L"), cutoff=1)
    return base.filter(ImageFilter.UnsharpMask(radius=1.2, percent=120, threshold=3))


def _sharpened(image: Any) -> Any:
    from PIL import ImageFilter

    return image.convert("L").filter(ImageFilter.UnsharpMask(radius=1.5, percent=160, threshold=2))


def _denoised(image: Any, noise: float) -> Any:
    from PIL import Image, ImageFilter

    try:
        import cv2
        import numpy as np

        gray = np.asarray(image.convert("L"))
        if noise >= 0.16:
            cleaned = cv2.fastNlMeansDenoising(gray, None, 7, 7, 21)
        else:
            cleaned = cv2.medianBlur(gray, 3)
        return Image.fromarray(cleaned)
    except Exception:  # noqa: BLE001 - PIL fallback
        return image.convert("L").filter(ImageFilter.MedianFilter(3))


def _binarized(image: Any, method: str) -> Any:
    from PIL import Image, ImageOps

    gray = ImageOps.grayscale(image)
    try:
        import cv2
        import numpy as np

        array = np.asarray(gray)
        if method == "adaptive":
            binary = cv2.adaptiveThreshold(
                array,
                255,
                cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                cv2.THRESH_BINARY,
                31,
                10,
            )
        else:
            _, binary = cv2.threshold(array, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        return Image.fromarray(binary)
    except Exception:  # noqa: BLE001 - PIL fallback
        threshold = _otsu_threshold(gray)
        return gray.point(lambda value: 255 if value > threshold else 0)


def _red_channel(image: Any) -> Any:
    return image.convert("RGB").split()[0]


def select_channel_image(path: Path, channel: str) -> Path | None:
    """Materialize a single-channel image; None keeps the original RGB path."""
    if channel not in {"gray", "red"}:
        return None
    try:
        from PIL import Image

        with Image.open(path) as opened:
            image = opened.convert("RGB")
    except Exception:  # noqa: BLE001 - channel selection is best-effort
        return None
    converted = image.convert("L") if channel == "gray" else _red_channel(image)
    temp = temp_image_path()
    converted.save(temp, format="PNG")
    return temp


def _candidate_score(metrics: dict[str, Any] | None) -> float:
    if not metrics:
        return -1.0
    return (
        float(metrics.get("sharpness") or 0) * 0.5
        + float(metrics.get("contrast") or 0) * 0.3
        + float(metrics.get("brightness") or 0) * 0.1
        - float(metrics.get("glare") or 0) * 1.5
        - max(0.0, float(metrics.get("shadow") or 0) - 0.2) * 0.4
        - float(metrics.get("noise") or 0) * 0.15
    )


def preprocess_image(
    path: Path, metrics: dict[str, Any] | None
) -> tuple[Path, list[Path], dict[str, Any]]:
    """Generate a bounded set of photometric candidates and keep the best.

    Only issue-driven strategies run, and a candidate must beat the original
    image score; otherwise the original path is returned unchanged.
    """
    try:
        from PIL import Image
    except Exception:  # noqa: BLE001 - Pillow is optional
        return path, [], {}

    try:
        with Image.open(path) as opened:
            image = opened.convert("RGB")
    except Exception:  # noqa: BLE001 - unreadable metadata should not fail OCR
        return path, [], {}

    issues = set((metrics or {}).get("issues") or [])
    shadow = float((metrics or {}).get("shadow") or 0)
    glare = float((metrics or {}).get("glare") or 0)
    noise = float((metrics or {}).get("noise") or 0)
    contrast = float((metrics or {}).get("contrast") or 0)
    sharpness = float((metrics or {}).get("sharpness") or 0)
    colorfulness = float((metrics or {}).get("colorfulness") or 0)

    candidates: list[dict[str, Any]] = [
        {"strategy": "original", "path": path, "metrics": metrics or {}}
    ]
    cleanup: list[Path] = []

    def add(strategy: str, candidate: Any) -> None:
        temp = temp_image_path()
        candidate.convert("RGB").save(temp, format="PNG")
        cleanup.append(temp)
        candidates.append(
            {"strategy": strategy, "path": temp, "metrics": image_metrics(candidate) or {}}
        )

    # Colour documents (stamps, highlights, tinted paper) can lose glyphs in a
    # plain grayscale conversion, so compare grayscale and red-only early.
    if colorfulness >= 0.08:
        add("grayscale", image.convert("L"))
        add("red_channel", _red_channel(image))
    # Glare cannot be reconstructed; do not binarize or sharpen a glare image
    # into fake glyphs. Contrast/shadow correction is still allowed.
    if shadow >= 0.3 or "shadow" in issues:
        add("shadow_corrected", _shadow_corrected(image))
    if contrast < 0.15 or issues & {"underexposed", "overexposed", "low_contrast"}:
        add("contrast", _contrast_enhanced(image))
    if noise >= 0.09 or "noisy" in issues:
        add("denoised", _denoised(image, noise))
    if glare < 0.05 and (sharpness < 0.2 or "blurry" in issues):
        add("sharpened", _sharpened(image))
    if glare < 0.02 and (contrast < 0.1 or "low_contrast" in issues):
        add("otsu", _binarized(image, "otsu"))
        add("adaptive", _binarized(image, "adaptive"))

    limit = env_int("OCR_PREPROCESS_MAX_CANDIDATES", 4)
    generated = candidates[1 : limit + 1]
    original = candidates[0]
    original_score = _candidate_score(original["metrics"])
    best = original
    best_score = original_score
    for candidate in generated:
        score = _candidate_score(candidate["metrics"])
        candidate["score"] = round(score, 4)
        if score > best_score:
            best = candidate
            best_score = score
    original["score"] = round(original_score, 4)

    minimum_gain = 0.01
    chosen = best if best is not original and best_score >= original_score + minimum_gain else original
    for candidate in cleanup:
        if candidate != chosen["path"]:
            candidate.unlink(missing_ok=True)
    cleanup = [chosen["path"]] if chosen is not original else []

    meta = {
        "strategy": chosen["strategy"],
        "gain": round(best_score - original_score, 4),
        "adopted": chosen is not original,
        "candidates": [
            {
                "strategy": candidate["strategy"],
                "score": candidate.get("score"),
                "sharpness": (candidate["metrics"] or {}).get("sharpness"),
                "contrast": (candidate["metrics"] or {}).get("contrast"),
                "glare": (candidate["metrics"] or {}).get("glare"),
                "shadow": (candidate["metrics"] or {}).get("shadow"),
                "noise": (candidate["metrics"] or {}).get("noise"),
            }
            for candidate in [original, *generated]
        ],
    }
    return chosen["path"], cleanup, meta


def detect_colored_regions(image: Any) -> list[dict[str, Any]]:
    """Cheap stamp/blue-ink hints; detects regions, not the ink itself."""
    try:
        import numpy as np
    except Exception:  # noqa: BLE001 - optional dependency
        return []
    try:
        width, height = image.size
        sample = image.convert("RGB")
        sample.thumbnail((1000, 1000))
        data = np.asarray(sample, dtype=np.int16)
        scale_x = width / sample.width if sample.width else 1.0
        scale_y = height / sample.height if sample.height else 1.0
        regions: list[dict[str, Any]] = []
        masks = (
            (
                (data[..., 0] - np.maximum(data[..., 1], data[..., 2]) > 40)
                & (data[..., 0] > 120),
                "stamp",
            ),
            (
                (data[..., 2] - np.maximum(data[..., 0], data[..., 1]) > 35)
                & (data[..., 2] > 90),
                "signature",
            ),
        )
        for mask, label in masks:
            ratio = float(mask.mean())
            if ratio < 0.0003 or ratio > 0.25:
                continue
            rows, columns = np.nonzero(mask)
            if rows.size < 20:
                continue
            left = float(columns.min()) * scale_x
            top = float(rows.min()) * scale_y
            right = float(columns.max() + 1) * scale_x
            bottom = float(rows.max() + 1) * scale_y
            if right - left < 20 or bottom - top < 20:
                continue
            aspect = (right - left) / max(bottom - top, 1)
            # Stamps are compact marks; a wide red banner or underline is not a
            # stamp. Blue signatures may be wide, so only constrain stamps.
            if label == "stamp" and (aspect < 0.25 or aspect > 4):
                continue
            regions.append(
                {
                    "label": label,
                    "bbox": box_bbox(left, top, right, bottom),
                    "ratio": round(ratio, 5),
                }
            )
        return regions
    except Exception:  # noqa: BLE001 - hints must never block OCR
        return []


def attach_colored_regions(
    blocks: list[dict[str, Any]],
    regions: list[dict[str, Any]],
    page_number: int,
    source: dict[str, Any],
) -> None:
    """Isolate stamp and blue-ink text from body paragraphs without rewriting it."""
    for region in regions:
        target = next(
            (
                block
                for block in blocks
                if block.get("type") == "paragraph"
                and _center_inside_bbox(block.get("bbox"), region["bbox"])
            ),
            None,
        )
        if target is not None:
            target["type"] = region["label"]
            target["meta"] = {
                **(target.get("meta") or {}),
                "region": region["label"],
                "regionRatio": region["ratio"],
                "regionSource": "color",
            }
            continue
        block = dir_block(
            f"p{page_number}-c{len(blocks)}",
            region["label"],
            "",
            region["bbox"],
            None,
            source,
            len(blocks),
        )
        block["meta"] = {
            "region": region["label"],
            "regionRatio": region["ratio"],
            "regionSource": "color",
        }
        blocks.append(block)


def _page_text_height(blocks: list[dict[str, Any]]) -> float | None:
    heights = [
        float(block["bbox"]["height"])
        for block in blocks
        if block.get("type") in {"paragraph", "heading"}
        and isinstance(block.get("bbox"), dict)
        and isinstance(block["bbox"].get("height"), (int, float))
        and block["bbox"]["height"] > 0
    ]
    if not heights:
        return None
    heights.sort()
    return heights[len(heights) // 2]


def detect_ink_gaps(
    image: Any,
    blocks: list[dict[str, Any]],
    page_number: int,
    colored_regions: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]] | None:
    """Report ink that no recognition box covers.

    Deliberately independent of recognition confidence: a missed line keeps its
    pixels even when the model says nothing suspicious about the text it did
    return. Every box the engine claimed is masked out, so what remains is either
    an unclaimed region or ink inside a table cell the engine left empty.

    An empty list means the page was analyzed and nothing was left unexplained;
    None means the analysis did not complete, and callers must not read that as
    clean.
    """
    try:
        import numpy as np
    except Exception:  # noqa: BLE001 - optional dependency
        return None
    max_gaps = env_int_zero_ok("OCR_INK_MAX_GAPS", DEFAULT_INK_MAX_GAPS)
    if max_gaps <= 0:
        return None
    try:
        width, height = image.size
        if width < 8 or height < 8:
            return None
        limit = max(64, env_int("OCR_INK_MAX_SIDE", DEFAULT_INK_MAX_SIDE))
        gray = image.convert("L")
        if max(width, height) > limit:
            gray.thumbnail((limit, limit))
        columns, rows = gray.size
        scale_x = columns / float(width)
        scale_y = rows / float(height)
        # `<=` keeps a two-level image (a scan with no antialiasing) from
        # producing an empty mask when Otsu lands on the ink value itself.
        mask = np.asarray(gray, dtype=np.uint8) <= _otsu_threshold(gray)
        ink_total = float(mask.sum())
        if ink_total < 24:
            # A page with no ink cannot have missed anything; that is an analyzed
            # clean page, not a failed analysis.
            return []
        coverage = np.zeros(mask.shape, dtype=bool)

        def box_of(bbox: Any) -> tuple[int, int, int, int] | None:
            if not isinstance(bbox, dict):
                return None
            # One pixel of slack on the far edges: the grid is a downscaled copy, so
            # truncating a box to floor() leaves an L-shaped sliver of "unexplained"
            # ink along every recognized line, which clusters into a full-size gap.
            left = max(0, int(math.floor(float(bbox.get("x") or 0.0) * scale_x)) - 1)
            top = max(0, int(math.floor(float(bbox.get("y") or 0.0) * scale_y)) - 1)
            right = min(
                columns,
                int(math.ceil((float(bbox.get("x") or 0.0) + float(bbox.get("width") or 0.0)) * scale_x)) + 1,
            )
            bottom = min(
                rows,
                int(math.ceil((float(bbox.get("y") or 0.0) + float(bbox.get("height") or 0.0)) * scale_y)) + 1,
            )
            if right <= left or bottom <= top:
                return None
            return left, top, right, bottom

        def cover(bbox: Any) -> None:
            box = box_of(bbox)
            if box:
                coverage[box[1] : box[3], box[0] : box[2]] = True

        def cover_frame(bbox: Any) -> None:
            """Mask a blank cell's ruled borders without masking its interior.

            An empty cell of a real grid is still drawn, and that line ink alone
            stays under the pixel and density floors. Anything actually written in
            the cell — handwriting, a stamp, a value the engine dropped — survives
            as uncovered ink, which is the second half of this detector's brief.
            Only ever sets coverage to True, so ordering against other boxes
            cannot un-explain ink someone else already claimed.
            """
            box = box_of(bbox)
            if not box:
                return
            left, top, right, bottom = box
            height = bottom - top
            breadth = right - left
            if height < 4 or breadth < 4:
                cover(bbox)
                return
            band_y = max(1, round(height * 0.25))
            band_x = max(1, round(breadth * 0.25))
            coverage[top:bottom, left : left + band_x] = True
            coverage[top:bottom, right - band_x : right] = True
            coverage[top : top + band_y, left:right] = True
            coverage[bottom - band_y : bottom, left:right] = True

        def cell_is_blank(value: Any) -> bool:
            text = str(value if value is not None else "").strip().replace("\u3000", "")
            return text == "" or bool(INK_BLANK_CELL_TEXT.match(text))

        # A table with per-cell coordinates explains the cells it read; a cell it
        # read as empty or as a placeholder explains only its borders.
        cell_tables: list[tuple[dict[str, Any], list[tuple[float, float, float, float]]]] = []
        for block in blocks:
            if block.get("type") != "table":
                cover(block.get("bbox"))
                continue
            entries = [
                entry
                for entry in (block.get("cells") or [])
                if isinstance(entry, dict) and isinstance(entry.get("bbox"), dict)
            ]
            if len(entries) < 2:
                # Without per-cell coordinates the grid cannot be attributed, so
                # treat the whole table as explained instead of inventing gaps.
                cover(block.get("bbox"))
                continue
            rects: list[tuple[float, float, float, float]] = []
            for entry in entries:
                bbox = entry["bbox"]
                if cell_is_blank(entry.get("text")):
                    cover_frame(bbox)
                else:
                    cover(bbox)
                    rects.append(
                        (
                            float(bbox.get("x") or 0.0),
                            float(bbox.get("y") or 0.0),
                            float(bbox.get("x") or 0.0) + float(bbox.get("width") or 0.0),
                            float(bbox.get("y") or 0.0) + float(bbox.get("height") or 0.0),
                        )
                    )
            cell_tables.append((block, rects))
        for region in colored_regions or []:
            cover(region.get("bbox"))

        free = mask & ~coverage
        if float(free.sum()) / ink_total > INK_UNSAFE_UNCOVERED_RATIO:
            # Most of the page is unexplained: the boxes do not line up with the
            # pixels. Aborted, not clean.
            return None

        def table_owner(
            left: float, top: float, right: float, bottom: float
        ) -> dict[str, Any] | None:
            """Attribute a gap to a table only inside a row the engine actually read.

            A `table_hint` grouping spans the whole page with one box per recognized
            line, so plain containment would relabel every missed paragraph as an
            empty cell and harden the risk. Requiring a filled cell in the same row
            band, that the gap does not sit on top of, keeps the cell signal for real
            grids.
            """
            gap_height = max(1.0, bottom - top)
            for block, cells in cell_tables:
                for cell_left, cell_top, cell_right, cell_bottom in cells:
                    band = min(gap_height, cell_bottom - cell_top)
                    if band <= 0 or min(bottom, cell_bottom) - max(top, cell_top) < 0.5 * band:
                        continue
                    if min(right, cell_right) - max(left, cell_left) > 0.3 * (right - left):
                        continue
                    return block
            return None

        # Ink is counted on the downscaled analysis grid, but the knob is page
        # pixels, so a smaller `OCR_INK_MAX_SIDE` cannot quietly raise the bar.
        min_pixels = max(8, env_int("OCR_INK_MIN_PIXELS", DEFAULT_INK_MIN_PIXELS))
        page_pixels_per_grid_pixel = 1.0 / max(1e-6, scale_x * scale_y)
        median_height = _page_text_height(blocks)
        step = max(6, int(round(min(columns, rows) / 60)))
        grid_rows = int(np.ceil(rows / step))
        grid_columns = int(np.ceil(columns / step))
        padded = np.zeros((grid_rows * step, grid_columns * step), dtype=bool)
        padded[:rows, :columns] = free
        counts = padded.reshape(grid_rows, step, grid_columns, step).sum(axis=(1, 3))
        active = counts >= max(3, int(step * step * 0.02))
        seen = np.zeros(active.shape, dtype=bool)
        # Every grid cell is stamped with the component that reached it, so a
        # component's ink is its own cells only: counting the bounding rectangle
        # would let an annulus and the figure inside it each report the other's
        # pixels, clear the pixel floor twice over and crop the same text into two
        # recovered blocks.
        labels = np.full(active.shape, -1, dtype=np.int32)
        extents: list[tuple[int, int, int, int]] = []
        for start_y in range(grid_rows):
            for start_x in range(grid_columns):
                if not active[start_y, start_x] or seen[start_y, start_x]:
                    continue
                index = len(extents)
                stack = [(start_y, start_x)]
                seen[start_y, start_x] = True
                top = bottom = start_y
                left = right = start_x
                while stack:
                    current_y, current_x = stack.pop()
                    labels[current_y, current_x] = index
                    top = min(top, current_y)
                    bottom = max(bottom, current_y)
                    left = min(left, current_x)
                    right = max(right, current_x)
                    for neighbor_y, neighbor_x in (
                        (current_y - 1, current_x),
                        (current_y + 1, current_x),
                        (current_y, current_x - 1),
                        (current_y, current_x + 1),
                    ):
                        if (
                            0 <= neighbor_y < grid_rows
                            and 0 <= neighbor_x < grid_columns
                            and active[neighbor_y, neighbor_x]
                            and not seen[neighbor_y, neighbor_x]
                        ):
                            seen[neighbor_y, neighbor_x] = True
                            stack.append((neighbor_y, neighbor_x))
                extents.append((top, left, bottom, right))
        flat_labels = labels.ravel()
        owned = flat_labels >= 0
        component_ink = np.bincount(
            flat_labels[owned], weights=counts.ravel()[owned], minlength=len(extents)
        )
        order = sorted(range(len(extents)), key=lambda index: component_ink[index], reverse=True)

        gaps: list[dict[str, Any]] = []
        for index in order[:40]:
            ys_list: list[Any] = []
            xs_list: list[Any] = []
            for cell_y, cell_x in zip(*np.nonzero(labels == index)):
                block_top = cell_y * step
                block_left = cell_x * step
                block = free[
                    block_top : min(rows, block_top + step),
                    block_left : min(columns, block_left + step),
                ]
                cell_ys, cell_xs = np.nonzero(block)
                if cell_ys.size:
                    ys_list.append(cell_ys + block_top)
                    xs_list.append(cell_xs + block_left)
            if not ys_list:
                continue
            ys = np.concatenate(ys_list)
            xs = np.concatenate(xs_list)
            ink_pixels = int(round(ys.size * page_pixels_per_grid_pixel))
            if ink_pixels < min_pixels:
                continue
            top = int(ys.min())
            bottom = int(ys.max()) + 1
            left = int(xs.min())
            right = int(xs.max()) + 1
            area = max(1, (bottom - top) * (right - left))
            density = float(ys.size) / area
            gap_left = left / scale_x
            gap_top = top / scale_y
            gap_width = (right - left) / scale_x
            gap_height = (bottom - top) / scale_y
            if density < 0.04 or gap_width < 8 or gap_height < 8:
                continue
            if median_height is not None and gap_height < median_height * 0.5:
                continue
            if gap_width * gap_height > 0.7 * width * height:
                continue
            owner = table_owner(gap_left, gap_top, gap_left + gap_width, gap_top + gap_height)
            gaps.append(
                {
                    "id": f"p{page_number}-g{len(gaps)}",
                    "kind": "cell" if owner else "region",
                    "blockId": owner.get("id") if owner else None,
                    "bbox": box_bbox(
                        gap_left, gap_top, gap_left + gap_width, gap_top + gap_height
                    ),
                    "pixels": ink_pixels,
                    "density": round(min(1.0, density), 4),
                }
            )
            if len(gaps) >= max_gaps:
                break
        return gaps
    except Exception:  # noqa: BLE001 - miss detection must never block OCR
        return None


def detect_page_ink_gaps(
    path: Path,
    blocks: list[dict[str, Any]],
    page_number: int,
    colored_regions: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]] | None:
    """Ink-gap detection for a finished page; None when the analysis did not run.

    The distinction matters downstream: an empty list means the page was checked
    and no ink was left unexplained, while None is the absence of evidence.
    """
    if not blocks or env_int_zero_ok("OCR_INK_MAX_GAPS", DEFAULT_INK_MAX_GAPS) <= 0:
        return None
    try:
        from PIL import Image

        with Image.open(path) as opened:
            return detect_ink_gaps(
                opened.convert("RGB"), blocks, page_number, colored_regions
            )
    except Exception:  # noqa: BLE001 - miss detection is optional
        return None


def _mrz_char_value(value: str) -> int:
    if value.isdigit():
        return int(value)
    if "A" <= value <= "Z":
        return ord(value) - ord("A") + 10
    return 0


def mrz_check_digit(value: str) -> str:
    """ICAO 9303 check digit; '<' counts as zero."""
    total = 0
    for index, char in enumerate(value.upper()):
        total += _mrz_char_value(char) * (7, 3, 1)[index % 3]
    return str(total % 10)


def _mrz_date(value: str, *, birth: bool) -> str | None:
    if not re.fullmatch(r"\d{6}", value):
        return None
    year = int(value[0:2])
    current = int(time.strftime("%y"))
    if birth:
        year += 1900 if year > current else 2000
    else:
        year += 1900 if year >= 70 else 2000
    month = int(value[2:4])
    day = int(value[4:6])
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    return f"{year:04d}-{month:02d}-{day:02d}"


def _mrz_checks(values: list[tuple[str, str, str]]) -> dict[str, Any]:
    failed = [name for name, value, expected in values if mrz_check_digit(value) != expected]
    return {"valid": not failed, "failed": failed}


def parse_mrz(text: str, page_number: int = 1) -> list[dict[str, Any]]:
    """Find TD1/TD2/TD3 MRZ bands in recognized text and verify check digits."""
    lines = [re.sub(r"[^A-Z0-9<]", "", line.strip().upper()) for line in text.splitlines()]
    lines = [line for line in lines if len(line) in (30, 36, 44) and "<" in line]
    found: list[dict[str, Any]] = []
    seen: set[str] = set()

    def add(entry: dict[str, Any]) -> None:
        key = f"{entry.get('text')}"
        if key in seen:
            return
        seen.add(key)
        entry["page"] = page_number
        found.append(entry)

    for index in range(len(lines) - 1):
        first, second = lines[index], lines[index + 1]
        if len(first) == 44 and len(second) == 44 and first[0] in "PIAC":
            number = second[0:9]
            birth = second[13:19]
            expiry = second[21:27]
            personal = second[28:42]
            checks = _mrz_checks(
                [
                    ("documentNumber", number, second[9]),
                    ("birthDate", birth, second[19]),
                    ("expiryDate", expiry, second[27]),
                    ("personalNumber", personal, second[42]),
                ]
            )
            names = first[5:44].split("<<", 1)
            add(
                {
                    "format": "mrz",
                    "text": f"{first}\n{second}",
                    "fields": {
                        "documentNumber": number.replace("<", "") or None,
                        "nationality": second[10:13].replace("<", "") or None,
                        "birthDate": _mrz_date(birth, birth=True),
                        "expiryDate": _mrz_date(expiry, birth=False),
                        "sex": second[20] if second[20] in "MF" else None,
                        "surname": names[0].replace("<", " ").strip() or None,
                        "givenNames": (names[1] if len(names) > 1 else "").replace("<", " ").strip()
                        or None,
                    },
                    "checks": checks,
                }
            )
        if len(first) == 36 and len(second) == 36 and first[0] in "IAC":
            number = second[0:9]
            birth = second[13:19]
            expiry = second[21:27]
            checks = _mrz_checks(
                [
                    ("documentNumber", number, second[9]),
                    ("birthDate", birth, second[19]),
                    ("expiryDate", expiry, second[27]),
                ]
            )
            names = first[5:36].split("<<", 1)
            add(
                {
                    "format": "mrz",
                    "text": f"{first}\n{second}",
                    "fields": {
                        "documentNumber": number.replace("<", "") or None,
                        "nationality": second[10:13].replace("<", "") or None,
                        "birthDate": _mrz_date(birth, birth=True),
                        "expiryDate": _mrz_date(expiry, birth=False),
                        "sex": second[20] if second[20] in "MF" else None,
                        "surname": names[0].replace("<", " ").strip() or None,
                        "givenNames": (names[1] if len(names) > 1 else "").replace("<", " ").strip()
                        or None,
                    },
                    "checks": checks,
                }
            )
    for index in range(len(lines) - 2):
        first, second, third = lines[index : index + 3]
        if not (
            len(first) == 30
            and len(second) == 30
            and len(third) == 30
            and first[0] in "IAC"
        ):
            continue
        number = first[5:14]
        birth = second[0:6]
        expiry = second[8:14]
        checks = _mrz_checks(
            [
                ("documentNumber", number, first[14]),
                ("birthDate", birth, second[6]),
                ("expiryDate", expiry, second[14]),
            ]
        )
        names = third.split("<<", 1)
        add(
            {
                "format": "mrz",
                "text": f"{first}\n{second}\n{third}",
                "fields": {
                    "documentNumber": number.replace("<", "") or None,
                    "nationality": second[15:18].replace("<", "") or None,
                    "birthDate": _mrz_date(birth, birth=True),
                    "expiryDate": _mrz_date(expiry, birth=False),
                    "sex": second[7] if second[7] in "MF" else None,
                    "surname": names[0].replace("<", " ").strip() or None,
                    "givenNames": (names[1] if len(names) > 1 else "").replace("<", " ").strip()
                    or None,
                },
                "checks": checks,
            }
        )
    return found


def _machine_bbox(points: Any) -> dict[str, float] | None:
    try:
        import numpy as np

        array = np.asarray(points, dtype="float32").reshape(-1, 2)
        if array.size == 0:
            return None
        left, top = array.min(axis=0)
        right, bottom = array.max(axis=0)
        return box_bbox(float(left), float(top), float(right), float(bottom))
    except Exception:  # noqa: BLE001 - bbox is advisory
        return None


def decode_machine_symbols(path: Path, page_number: int = 1) -> list[dict[str, Any]]:
    """Decode QR/barcodes with whatever optional decoder is installed.

    OpenCV reads QR codes; pyzbar additionally reads common 1D barcodes. The
    payload is evidence only: when neither library exists the page still gets
    MRZ evidence parsed from text.
    """
    if env_int_zero_ok("OCR_MACHINE_READABLE", 1) <= 0:
        return []
    limit = env_int("OCR_MACHINE_READABLE_MAX_SYMBOLS", 8)
    found: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()

    def add(format_name: str, payload: str, bbox: dict[str, float] | None) -> None:
        text = payload.strip()
        if not text or len(text) > 4000:
            return
        key = (format_name, text)
        if key in seen or len(found) >= limit:
            return
        seen.add(key)
        found.append({"format": format_name, "text": text, "page": page_number, "bbox": bbox})

    try:
        import cv2

        image = cv2.imread(str(path))
        if image is not None:
            detector = cv2.QRCodeDetector()
            multi = getattr(detector, "detectAndDecodeMulti", None)
            decoded: Any = None
            points: Any = None
            if callable(multi):
                ok, decoded, points, _ = multi(image)
                if not ok:
                    decoded = None
            else:
                payload, qr_points, _ = detector.detectAndDecode(image)
                if payload:
                    decoded = [payload]
                    points = [qr_points] if qr_points is not None else None
            if decoded:
                for index, payload in enumerate(decoded):
                    if not payload:
                        continue
                    bbox = (
                        _machine_bbox(points[index])
                        if points is not None and index < len(points)
                        else None
                    )
                    add("qr", str(payload), bbox)
    except Exception:  # noqa: BLE001 - decoder is optional
        pass

    try:
        from PIL import Image
        from pyzbar.pyzbar import decode as zbar_decode

        with Image.open(path) as opened:
            for symbol in zbar_decode(opened):
                try:
                    payload = symbol.data.decode("utf-8", "replace")
                except Exception:  # noqa: BLE001 - corrupt payload
                    continue
                rect = symbol.rect
                bbox = box_bbox(
                    float(rect.left),
                    float(rect.top),
                    float(rect.left + rect.width),
                    float(rect.top + rect.height),
                )
                add("qr" if symbol.type == "QRCODE" else "barcode", payload, bbox)
    except Exception:  # noqa: BLE001 - decoder is optional
        pass
    return found


def machine_readable_for_page(
    path: Path, text: str, page_number: int
) -> list[dict[str, Any]]:
    entries = decode_machine_symbols(path, page_number)
    entries.extend(parse_mrz(text, page_number))
    return entries


def attach_machine_readable(page: dict[str, Any], entries: list[dict[str, Any]]) -> None:
    if not entries:
        return
    meta = dict(page.get("meta") or {})
    existing = list(meta.get("machineReadable") or [])
    seen = {(item.get("format"), item.get("text")) for item in existing}
    for entry in entries:
        key = (entry.get("format"), entry.get("text"))
        if key in seen:
            continue
        seen.add(key)
        existing.append(entry)
    meta["machineReadable"] = existing[:16]
    page["meta"] = meta


def _order_quad(points: Any) -> Any:
    import numpy as np

    ordered = np.zeros((4, 2), dtype="float32")
    sums = points.sum(axis=1)
    diffs = np.diff(points, axis=1).reshape(-1)
    ordered[0] = points[np.argmin(sums)]
    ordered[2] = points[np.argmax(sums)]
    ordered[1] = points[np.argmin(diffs)]
    ordered[3] = points[np.argmax(diffs)]
    return ordered


def correct_geometry(
    path: Path, metrics: dict[str, Any] | None
) -> tuple[Path | None, list[Path], dict[str, Any] | None]:
    """Optional page-boundary perspective/deskew. Off unless enabled.

    When applied, the caller maps DIR bboxes back through ``inverse`` so
    coordinates stay in the original image space.
    """
    if os.environ.get("OCR_PREPROCESS_GEOMETRY", "").strip().lower() not in {
        "1",
        "true",
        "yes",
    }:
        return None, [], None
    try:
        import cv2
        import numpy as np
    except Exception:  # noqa: BLE001 - optional dependency
        return None, [], None
    try:
        image = cv2.imread(str(path), cv2.IMREAD_COLOR)
        if image is None:
            return None, [], None
        height, width = image.shape[:2]
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        blurred = cv2.GaussianBlur(gray, (5, 5), 0)
        edges = cv2.Canny(blurred, 50, 150)
        contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
        quad = None
        for contour in sorted(contours, key=cv2.contourArea, reverse=True)[:5]:
            area = cv2.contourArea(contour)
            if area < 0.3 * width * height or area > 0.98 * width * height:
                continue
            approx = cv2.approxPolyDP(
                contour, 0.02 * cv2.arcLength(contour, True), True
            )
            if len(approx) != 4:
                continue
            quad = _order_quad(approx.reshape(4, 2).astype("float32"))
            break

        matrix = None
        warped = None
        strategy = None
        if quad is not None:
            target_width = int(
                max(
                    np.linalg.norm(quad[1] - quad[0]),
                    np.linalg.norm(quad[2] - quad[3]),
                )
            )
            target_height = int(
                max(
                    np.linalg.norm(quad[3] - quad[0]),
                    np.linalg.norm(quad[2] - quad[1]),
                )
            )
            if target_width > 50 and target_height > 50:
                destination = np.array(
                    [
                        [0, 0],
                        [target_width - 1, 0],
                        [target_width - 1, target_height - 1],
                        [0, target_height - 1],
                    ],
                    dtype="float32",
                )
                matrix = cv2.getPerspectiveTransform(quad, destination)
                warped = cv2.warpPerspective(image, matrix, (target_width, target_height))
                strategy = "perspective"
        if warped is None:
            skew = (metrics or {}).get("skew")
            if isinstance(skew, (int, float)) and 0.5 <= abs(float(skew)) <= 12:
                rotation = cv2.getRotationMatrix2D(
                    (width / 2, height / 2), -float(skew), 1.0
                )
                cosine = abs(rotation[0, 0])
                sine = abs(rotation[0, 1])
                new_width = int(height * sine + width * cosine)
                new_height = int(height * cosine + width * sine)
                rotation[0, 2] += new_width / 2 - width / 2
                rotation[1, 2] += new_height / 2 - height / 2
                warped = cv2.warpAffine(
                    image,
                    rotation,
                    (new_width, new_height),
                    flags=cv2.INTER_CUBIC,
                    borderValue=(255, 255, 255),
                )
                matrix = np.vstack([rotation, [0, 0, 1]])
                strategy = "deskew"
        if warped is None or matrix is None:
            return None, [], None
        temp = temp_image_path()
        cv2.imwrite(str(temp), warped)
        inverse = np.linalg.inv(matrix).tolist()
        return temp, [temp], {
            "strategy": strategy,
            "matrix": matrix.tolist(),
            "inverse": inverse,
            "originalSize": [width, height],
            "warpedSize": [warped.shape[1], warped.shape[0]],
        }
    except Exception:  # noqa: BLE001 - geometry is best-effort
        return None, [], None


def _map_point(inverse: list[list[float]], x: float, y: float) -> tuple[float, float] | None:
    denominator = inverse[2][0] * x + inverse[2][1] * y + inverse[2][2]
    if abs(denominator) < 1e-9:
        return None
    return (
        (inverse[0][0] * x + inverse[0][1] * y + inverse[0][2]) / denominator,
        (inverse[1][0] * x + inverse[1][1] * y + inverse[1][2]) / denominator,
    )


def map_bbox_back(
    bbox: dict[str, float] | None, inverse: list[list[float]], width: float, height: float
) -> dict[str, float] | None:
    if not isinstance(bbox, dict):
        return bbox
    x = float(bbox.get("x") or 0)
    y = float(bbox.get("y") or 0)
    box_width = float(bbox.get("width") or 0)
    box_height = float(bbox.get("height") or 0)
    corners = [(x, y), (x + box_width, y), (x, y + box_height), (x + box_width, y + box_height)]
    mapped = [_map_point(inverse, point[0], point[1]) for point in corners]
    if any(point is None for point in mapped):
        return bbox
    xs = [point[0] for point in mapped if point is not None]
    ys = [point[1] for point in mapped if point is not None]
    left = max(0.0, min(xs))
    top = max(0.0, min(ys))
    right = min(width, max(xs))
    bottom = min(height, max(ys))
    return box_bbox(left, top, right, bottom)


def map_blocks_back(
    blocks: list[dict[str, Any]],
    inverse: list[list[float]],
    width: float,
    height: float,
) -> list[dict[str, Any]]:
    for block in blocks:
        block["bbox"] = map_bbox_back(block.get("bbox"), inverse, width, height)
        for cell in block.get("cells") or []:
            if isinstance(cell, dict):
                cell["bbox"] = map_bbox_back(cell.get("bbox"), inverse, width, height)
    return blocks


def read_pdfplumber_pages(
    path: Path, max_pages: int | None = None
) -> list[dict[str, Any]] | None:
    """Native PDF text, line boxes and table grids; ``None`` when unavailable."""
    try:
        import pdfplumber
    except Exception:  # noqa: BLE001 - optional native extractor
        return None
    try:
        pages: list[dict[str, Any]] = []
        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages):
                if max_pages is not None and index >= max_pages:
                    break
                try:
                    text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                except Exception:  # noqa: BLE001 - broken text layer on one page
                    text = ""
                lines: list[dict[str, Any]] = []
                try:
                    for line in page.extract_text_lines(x_tolerance=2, y_tolerance=3) or []:
                        chars = line.get("chars") or []
                        sizes = [float(char.get("size", 0)) for char in chars]
                        sizes = [size for size in sizes if size > 0]
                        lines.append(
                            {
                                "text": str(line.get("text") or ""),
                                "x0": float(line.get("x0") or 0),
                                "top": float(line.get("top") or 0),
                                "x1": float(line.get("x1") or 0),
                                "bottom": float(line.get("bottom") or 0),
                                "size": sum(sizes) / len(sizes) if sizes else None,
                            }
                        )
                except Exception:  # noqa: BLE001 - line bboxes are optional
                    lines = []
                tables: list[dict[str, Any]] = []
                try:
                    for table in page.find_tables() or []:
                        rows = _clean_rows(table.extract() or [])
                        if len(rows) < 2:
                            continue
                        bbox = getattr(table, "bbox", None) or (0, 0, 0, 0)
                        tables.append(
                            {
                                "rows": rows,
                                "bbox": box_bbox(
                                    float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3])
                                ),
                            }
                        )
                except Exception:  # noqa: BLE001 - table detection is best effort
                    tables = []
                images: list[dict[str, Any]] = []
                try:
                    for image in getattr(page, "images", []) or []:
                        x0 = float(image.get("x0") or 0)
                        top = float(image.get("top") or 0)
                        x1 = float(image.get("x1") or 0)
                        bottom = float(image.get("bottom") or 0)
                        if x1 <= x0 or bottom <= top:
                            continue
                        images.append(
                            {
                                "bbox": box_bbox(x0, top, x1, bottom),
                                "width": x1 - x0,
                                "height": bottom - top,
                            }
                        )
                except Exception:  # noqa: BLE001 - image boxes are optional
                    images = []
                trust, trust_reasons = text_layer_trust(text)
                pages.append(
                    {
                        "text": text,
                        "width": float(getattr(page, "width", 0) or 0) or None,
                        "height": float(getattr(page, "height", 0) or 0) or None,
                        "trust": trust,
                        "trustReasons": trust_reasons,
                        "lines": lines,
                        "tables": tables,
                        "images": images,
                    }
                )
        return pages
    except Exception as cause:  # noqa: BLE001 - encrypted or malformed PDF
        print(f"pdfplumber unavailable for this PDF ({cause})", file=sys.stderr)
        return None


def _center_inside_bbox(
    inner: dict[str, float] | None, outer: dict[str, float] | None
) -> bool:
    if not inner or not outer:
        return False
    center_x = inner["x"] + inner["width"] / 2
    center_y = inner["y"] + inner["height"] / 2
    return (
        outer["x"] <= center_x <= outer["x"] + outer["width"]
        and outer["y"] <= center_y <= outer["y"] + outer["height"]
    )


def _line_inside_bbox(line: dict[str, Any], bbox: dict[str, float] | None) -> bool:
    if not bbox:
        return False
    inner = box_bbox(
        float(line.get("x0") or 0),
        float(line.get("top") or 0),
        float(line.get("x1") or 0),
        float(line.get("bottom") or 0),
    )
    return _center_inside_bbox(inner, bbox)


def native_pdf_dir_page(
    page: dict[str, Any], number: int, engine: str = "pdfplumber"
) -> dict[str, Any]:
    lines = [line for line in page.get("lines") or [] if line.get("text", "").strip()]
    sizes = [line["size"] for line in lines if line.get("size")]
    median = sorted(sizes)[len(sizes) // 2] if sizes else None
    source = dir_source(engine, None, "native")

    table_items: list[tuple[dict[str, float] | None, list[list[str]]]] = []
    for table in page.get("tables") or []:
        rows = table.get("rows") or []
        if len(rows) < 2:
            continue
        table_items.append((table.get("bbox"), rows))

    # pdfplumber returns table cell text again as text lines; drop those lines
    # so the same content is not reported twice in the DIR.
    body_lines = [
        line
        for line in lines
        if not any(_line_inside_bbox(line, bbox) for bbox, _ in table_items)
    ]
    entries: list[tuple[float, str, Any]] = [
        (float(line.get("top") or 0), "line", line) for line in body_lines
    ]
    for bbox, rows in table_items:
        top = float(bbox["y"]) if bbox else float("inf")
        entries.append((top, "table", (bbox, rows)))
    entries.sort(key=lambda item: item[0])

    blocks: list[dict[str, Any]] = []
    line_count = 0
    table_count = 0
    for order, (_, kind, payload) in enumerate(entries):
        if kind == "line":
            line = payload
            text = line["text"].strip()
            block_type = (
                "heading"
                if median and line.get("size") and line["size"] > median * 1.15 and len(text) <= 120
                else "paragraph"
            )
            blocks.append(
                dir_block(
                    f"p{number}-b{line_count}",
                    block_type,
                    text,
                    box_bbox(line["x0"], line["top"], line["x1"], line["bottom"]),
                    None,
                    source,
                    order,
                )
            )
            line_count += 1
        else:
            bbox, rows = payload
            blocks.append(
                dir_block(
                    f"p{number}-t{table_count}",
                    "table",
                    "\n".join(" | ".join(row) for row in rows),
                    bbox,
                    None,
                    source,
                    order,
                    rows=rows,
                )
            )
            table_count += 1

    if not blocks and page.get("text", "").strip():
        blocks.append(
            dir_block(f"p{number}-b0", "paragraph", page["text"], None, None, source, 0)
        )
    result = dir_page(number, page.get("width"), page.get("height"), blocks, source)
    attach_machine_readable(result, parse_mrz(str(page.get("text") or ""), number))
    return result


class PaddleEngine:
    name = "paddle"

    def __init__(self) -> None:
        from paddleocr import PaddleOCR

        det = os.environ.get("OCR_PADDLE_DET_MODEL", "PP-OCRv6_medium_det").strip()
        rec = os.environ.get("OCR_PADDLE_REC_MODEL", "PP-OCRv6_medium_rec").strip()
        self.det_model: str | None = det or None
        self.rec_model: str | None = rec or None
        self._table_pipeline: Any = None
        self._enhanced_ocr: Any = None
        self._ocr_by_lang: dict[str, Any] = {}
        self._enhanced_by_lang: dict[str, Any] = {}
        self._profiles = parse_model_profiles()
        self._profile_ocr: dict[tuple[str, bool], Any] = {}
        self._model_sha256: dict[str, str | None] = {}
        common = {
            "use_doc_unwarping": False,
            "use_textline_orientation": False,
            "enable_mkldnn": False,
        }
        attempts = [
            (det, rec, True),
            ("PP-OCRv5_server_det", "PP-OCRv5_server_rec", True),
            (None, None, True),
            (None, None, False),
        ]
        for det_name, rec_name, orientation in attempts:
            try:
                kwargs: dict[str, Any] = dict(common)
                kwargs["use_doc_orientation_classify"] = orientation
                if det_name:
                    kwargs["text_detection_model_name"] = det_name
                if rec_name:
                    kwargs["text_recognition_model_name"] = rec_name
                self._ocr = PaddleOCR(**kwargs)
                self.model = model_label(det_name or "default", rec_name or "default")
                self.det_model = det_name
                self.rec_model = rec_name
                break
            except Exception as cause:  # noqa: BLE001 - try the next model set
                print(
                    f"models {det_name or 'default'} "
                    f"(orientation={orientation}) unavailable ({cause}); trying fallback",
                    file=sys.stderr,
                )
        else:  # pragma: no cover - PaddleOCR itself failed to initialise
            raise RuntimeError("no usable PaddleOCR model set")

    def _enhanced_instance(self) -> Any:
        if self._enhanced_ocr is not None:
            return self._enhanced_ocr
        try:
            from paddleocr import PaddleOCR

            kwargs: dict[str, Any] = {
                "use_doc_orientation_classify": True,
                "use_doc_unwarping": True,
                "use_textline_orientation": True,
                "enable_mkldnn": False,
            }
            if self.det_model:
                kwargs["text_detection_model_name"] = self.det_model
            if self.rec_model:
                kwargs["text_recognition_model_name"] = self.rec_model
            self._enhanced_ocr = PaddleOCR(**kwargs)
        except Exception as cause:  # noqa: BLE001 - enhanced mode is a best-effort recovery
            print(f"enhanced pipeline unavailable ({cause}); using base OCR", file=sys.stderr)
            self._enhanced_ocr = self._ocr
        return self._enhanced_ocr

    def _profile_instance(self, profile: str | None, lang: str | None, enhanced: bool) -> Any:
        """Named model pair from ``OCR_MODEL_PROFILES``.

        Unknown profiles are an explicit error: model comparison and canary
        runs must fail loudly rather than measure the default model twice.
        """
        key = (profile or "").strip()
        if not key or key == "default":
            return self._language_instance(lang, enhanced)
        spec = self._profiles.get(key)
        if spec is None:
            raise ValueError(f"unknown model profile {key}")
        cache_key = (key, enhanced)
        if cache_key in self._profile_ocr:
            return self._profile_ocr[cache_key]
        try:
            from paddleocr import PaddleOCR

            kwargs: dict[str, Any] = {
                "use_doc_orientation_classify": enhanced,
                "use_doc_unwarping": enhanced,
                "use_textline_orientation": enhanced,
                "enable_mkldnn": False,
            }
            if spec.get("det"):
                kwargs["text_detection_model_name"] = spec["det"]
            if spec.get("rec"):
                kwargs["text_recognition_model_name"] = spec["rec"]
            elif lang and lang != "ch":
                kwargs["lang"] = lang
            instance = PaddleOCR(**kwargs)
        except Exception as cause:  # noqa: BLE001 - explicit profile failed
            raise ValueError(f"model profile {key} unavailable: {cause}") from cause
        self._profile_ocr[cache_key] = instance
        return instance

    def _language_instance(self, lang: str | None, enhanced: bool) -> Any:
        """Lazily build a script-specific recognizer; fall back to the default.

        Language hints come from the explainable router. We never download a
        different model silently on an ordinary path: a creation failure is
        logged and the default model is used instead.
        """
        key = (lang or "").strip().lower()
        if not key or key == "ch":
            return self._enhanced_instance() if enhanced else self._ocr
        cache = self._enhanced_by_lang if enhanced else self._ocr_by_lang
        if key in cache:
            return cache[key]
        try:
            from paddleocr import PaddleOCR

            kwargs: dict[str, Any] = {
                "lang": key,
                "use_doc_orientation_classify": enhanced,
                "use_doc_unwarping": enhanced,
                "use_textline_orientation": enhanced,
                "enable_mkldnn": False,
            }
            instance = PaddleOCR(**kwargs)
            cache[key] = instance
            return instance
        except Exception as cause:  # noqa: BLE001 - keep the default recognizer
            print(f"language {key} model unavailable ({cause}); using default", file=sys.stderr)
            fallback = self._enhanced_instance() if enhanced else self._ocr
            cache[key] = fallback
            return fallback

    def _model_for_lang(self, lang: str | None) -> str:
        key = (lang or "").strip().lower()
        return f"{self.model}/lang-{key}" if key and key != "ch" else self.model

    def _model_for(self, lang: str | None, profile: str | None = None) -> str:
        key = (profile or "").strip()
        if not key or key == "default":
            return self._model_for_lang(lang)
        spec = self._profiles.get(key) or {}
        label = model_label(spec.get("det") or "default", spec.get("rec") or "default")
        suffix = [key]
        lang_key = (lang or "").strip().lower()
        if lang_key and lang_key != "ch" and not spec.get("rec"):
            suffix.append(f"lang-{lang_key}")
        return "/".join([label, *suffix])

    def model_sha256(self, profile: str | None = None) -> str | None:
        key = (profile or "").strip()
        if key in self._model_sha256:
            return self._model_sha256[key]
        spec = self._profiles.get(key) if key and key != "default" else None
        names = [name for name in ((spec or {}).get("det"), (spec or {}).get("rec")) if name]
        if not names:
            names = [name for name in (self.det_model, self.rec_model) if name]
        value = model_files_sha256(names)
        self._model_sha256[key] = value
        return value

    def extract(
        self,
        path: Path,
        mode: str = "auto",
        dpi: int | None = None,
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> Extraction:
        if path.suffix.lower() == ".pdf":
            return self._extract_pdf(path, dpi, mode, lang, model, channel)
        validate_image_budget(path)
        return self._extract_images(path, mode, lang, model, channel)

    def _extract_images(
        self,
        path: Path,
        mode: str,
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> Extraction:
        frames, cleanup, truncated = prepare_image_frames(path)
        try:
            extractions: list[PageExtraction] = []
            scores: list[float] = []
            labels: set[str] = set()
            for index, frame in enumerate(frames, 1):
                page = self._page_from_image(frame, mode, index, lang, model, channel)
                extractions.append(page)
                labels.add(page.model)
                if page.confidence is not None:
                    scores.append(page.confidence)
            texts = [page.text for page in extractions]
            dir_pages = [page.page for page in extractions]
            text = "\n\n".join(page for page in texts if page.strip()).strip()
            confidence = sum(scores) / len(scores) if scores else None
            base_model = self._model_for(lang, model)
            label = labels.pop() if len(labels) == 1 else base_model
            meta: dict[str, Any] = {"mode": mode, "frames": len(frames)}
            if channel != "rgb":
                meta["channel"] = channel
            if lang:
                meta["lang"] = lang
            preprocess_pages = [page.preprocess for page in extractions if page.preprocess]
            if preprocess_pages:
                meta["preprocess"] = {"pages": preprocess_pages}
            document = dir_document("visual", dir_pages, meta)
            return Extraction(
                normalize_ocr_text(text),
                confidence,
                label,
                base_model,
                len(frames),
                truncated,
                "visual",
                document,
            )
        finally:
            for temp in cleanup:
                temp.unlink(missing_ok=True)

    def _page_from_image(
        self,
        path: Path,
        mode: str,
        page_number: int,
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> PageExtraction:
        cleanup: list[Path] = []
        if channel != "rgb":
            channel_path = select_channel_image(path, channel)
            if channel_path is not None:
                path = channel_path
                cleanup.append(channel_path)
        if mode in {"table", "structure"}:
            structured = self._structure_page(path, page_number, mode)
            if structured is not None:
                for temp in cleanup:
                    temp.unlink(missing_ok=True)
                return structured
        enhanced = mode == "enhanced"
        ocr = self._profile_instance(model, lang, enhanced)
        ocr_path = path
        base_path = path
        preprocess: dict[str, Any] | None = None
        geometry: dict[str, Any] | None = None
        if enhanced:
            metrics: dict[str, Any] | None = None
            try:
                from PIL import Image

                with Image.open(path) as opened:
                    metrics = image_metrics(opened.convert("RGB"))
            except Exception:  # noqa: BLE001 - enhanced preprocessing is best-effort
                metrics = None
            if metrics is not None:
                geometry_path, geometry_cleanup, geometry = correct_geometry(path, metrics)
                base_path = geometry_path or path
                cleanup.extend(geometry_cleanup)
                base_metrics = metrics
                if geometry is not None:
                    try:
                        from PIL import Image

                        with Image.open(base_path) as opened:
                            base_metrics = image_metrics(opened.convert("RGB")) or metrics
                    except Exception:  # noqa: BLE001 - keep original metrics
                        base_metrics = metrics
                ocr_path, candidate_cleanup, photometric = preprocess_image(
                    base_path, base_metrics
                )
                cleanup.extend(candidate_cleanup)
                preprocess = {
                    "page": page_number,
                    "geometry": geometry,
                    "photometric": photometric,
                    "source": str(path),
                    "applied": str(ocr_path),
                }
        color_regions: list[dict[str, Any]] = []
        try:
            from PIL import Image

            with Image.open(path) as opened:
                color_regions = detect_colored_regions(opened.convert("RGB"))
        except Exception:  # noqa: BLE001 - color hints are best-effort
            color_regions = []
        base_model = self._model_for(lang, model)
        label = f"{base_model}+enhanced" if enhanced else base_model

        def recognize(target: Path) -> tuple[str, float | None, list[dict[str, Any]], float | None, float | None]:
            results = list(ocr.predict(input=str(target)))
            if not results:
                raise RuntimeError("PaddleOCR returned no result")
            grid = None if enhanced else detect_table_grid(target)
            return paddle_page_blocks(
                results[0],
                page_number,
                "visual",
                base_model,
                label,
                table_hint=grid is not None,
                table_bbox=grid,
            )

        try:
            parsed = recognize(ocr_path)
            if enhanced and preprocess and ocr_path != base_path:
                # Image metrics propose a candidate; a signal other than the
                # candidate's own confidence has to accept it. Numeric tokens
                # the original preserved must not disappear without a conflict.
                baseline = recognize(base_path)
                selection = photometric_selection(parsed, baseline)
                preprocess["ocrSelection"] = selection
                if selection["selected"] == "original":
                    parsed = baseline
            text, confidence, blocks, width, height = parsed
            if geometry is not None:
                original_width, original_height = geometry["originalSize"]
                blocks = map_blocks_back(
                    blocks, geometry["inverse"], original_width, original_height
                )
                width = float(original_width)
                height = float(original_height)
            if color_regions:
                attach_colored_regions(
                    blocks,
                    color_regions,
                    page_number,
                    dir_source(base_model, label, "visual"),
                )
            ink_gaps = detect_page_ink_gaps(path, blocks, page_number, color_regions)
            meta: dict[str, Any] = {}
            if ink_gaps is not None:
                meta["inkGaps"] = ink_gaps
            page = dir_page(
                page_number,
                width,
                height,
                blocks,
                dir_source(base_model, label, "visual"),
                meta or None,
            )
            attach_machine_readable(page, machine_readable_for_page(path, text, page_number))
            return PageExtraction(
                normalize_ocr_text(text), confidence, page, label, preprocess
            )
        finally:
            for temp in cleanup:
                temp.unlink(missing_ok=True)

    def _table_pipeline_instance(self) -> Any:
        if self._table_pipeline is not None:
            return self._table_pipeline
        from paddleocr import TableRecognitionPipelineV2

        common: dict[str, Any] = {
            "use_doc_orientation_classify": True,
            "use_doc_unwarping": False,
            "use_layout_detection": True,
            "use_ocr_model": True,
            "enable_mkldnn": False,
        }
        if self.det_model:
            common["text_detection_model_name"] = self.det_model
        if self.rec_model:
            common["text_recognition_model_name"] = self.rec_model
        pinned = {
            **common,
            "layout_detection_model_name": "PP-DocLayout_plus-L",
            "table_classification_model_name": "PP-LCNet_x1_0_table_cls",
            "wired_table_structure_recognition_model_name": "SLANeXt_wired",
            "wired_table_cells_detection_model_name": "RT-DETR-L_wired_table_cell_det",
            "wireless_table_structure_recognition_model_name": "SLANet_plus",
            "wireless_table_cells_detection_model_name": "RT-DETR-L_wireless_table_cell_det",
        }
        last_error: Exception | None = None
        for kwargs in (pinned, common):
            try:
                self._table_pipeline = TableRecognitionPipelineV2(**kwargs)
                return self._table_pipeline
            except Exception as cause:  # noqa: BLE001 - try the default model set
                last_error = cause
                print(f"table pipeline model set unavailable ({cause})", file=sys.stderr)
        raise RequestError(422, f"表格结构化模型不可用：{str(last_error)[:200]}")

    def _structure_page(
        self, path: Path, page_number: int, mode: str
    ) -> PageExtraction | None:
        label = "PP-StructureV3-table" if mode == "table" else "PP-StructureV3"
        try:
            pipeline = self._table_pipeline_instance()
            results = list(pipeline.predict(input=str(path)))
        except RequestError:
            raise
        except Exception as cause:  # noqa: BLE001 - fall back to line-aware OCR
            print(f"structure pipeline failed ({cause}); falling back", file=sys.stderr)
            return None

        table_rows: list[list[str]] = []
        table_blocks: list[dict[str, Any]] = []
        table_boxes: list[dict[str, float]] = []
        overall_confidence: float | None = None
        overall_blocks: list[dict[str, Any]] = []
        layout_regions: list[dict[str, Any]] = []
        width: float | None = None
        height: float | None = None
        source = dir_source(self.model, label, "visual")

        for result in results:
            layout = result.get("layout_det_res") if hasattr(result, "get") else None
            boxes = layout.get("boxes") if hasattr(layout, "get") else None
            if boxes is None:
                boxes = getattr(layout, "boxes", None)
            for entry in boxes or []:
                if not hasattr(entry, "get"):
                    continue
                region_label = str(entry.get("label") or "").strip().lower()
                bbox = bbox_from_value(
                    entry.get("coordinate") or entry.get("bbox") or entry.get("box")
                )
                if not region_label or not bbox:
                    continue
                layout_regions.append(
                    {
                        "label": region_label,
                        "bbox": bbox,
                        "confidence": float(entry.get("score") or 0.0) or None,
                    }
                )
                width = max(width or 0.0, bbox["x"] + bbox["width"])
                height = max(height or 0.0, bbox["y"] + bbox["height"])
            tables = result.get("table_res_list") or []
            if not isinstance(tables, list):
                tables = [tables]
            for table in tables:
                if not hasattr(table, "get"):
                    continue
                pred_html = table.get("pred_html")
                grid = table_html_grid(str(pred_html)) if pred_html else []
                if not grid:
                    continue
                table_rows.extend(grid)
                bbox = bbox_from_value(
                    table.get("table_bbox")
                    or table.get("bbox")
                    or table.get("table_region")
                )
                table_blocks.append(
                    dir_block(
                        f"p{page_number}-t{len(table_blocks)}",
                        "table",
                        "\n".join(" | ".join(row) for row in grid),
                        bbox,
                        None,
                        source,
                        len(table_blocks),
                        rows=grid,
                        cells=table_pred_cells(grid, table),
                    )
                )
                if bbox:
                    table_boxes.append(bbox)
                    width = max(width or 0.0, bbox["x"] + bbox["width"])
                    height = max(height or 0.0, bbox["y"] + bbox["height"])
            overall = result.get("overall_ocr_res") or {}
            if hasattr(overall, "get") and overall.get("rec_texts"):
                _, _, blocks, page_width, page_height = paddle_page_blocks(
                    overall, page_number, "visual", self.model, label
                )
                overall_blocks.extend(blocks)
                width = max(width or 0.0, page_width or 0.0) or width
                height = max(height or 0.0, page_height or 0.0) or height

        # The layout pipeline reports table cells again through the overall OCR
        # result; drop those blocks so text, blocks and reading order are unique.
        if table_boxes:
            overall_blocks = [
                block
                for block in overall_blocks
                if not any(
                    _center_inside_bbox(block.get("bbox"), bbox) for bbox in table_boxes
                )
            ]
        overall_text = "\n".join(
            block["text"].strip() for block in overall_blocks if block["text"].strip()
        )
        overall_scores = [
            block["confidence"]
            for block in overall_blocks
            if block.get("confidence") is not None
        ]
        if overall_scores:
            overall_confidence = sum(overall_scores) / len(overall_scores)

        if not table_rows and not overall_text.strip():
            return None
        table_text = "\n".join(" | ".join(row) for row in table_rows)
        text = "\n\n".join(part for part in (table_text, overall_text) if part.strip())
        blocks = sorted(
            table_blocks + overall_blocks,
            key=lambda block: (
                (block.get("bbox") or {}).get("y", float("inf")),
                (block.get("bbox") or {}).get("x", 0.0),
            ),
        )
        for region in layout_regions:
            mapped = LAYOUT_LABELS.get(region["label"])
            if not mapped:
                continue
            block_type, region_name = mapped
            target = next(
                (
                    block
                    for block in blocks
                    if _center_inside_bbox(block.get("bbox"), region["bbox"])
                ),
                None,
            )
            if target is not None:
                target["meta"] = {
                    **(target.get("meta") or {}),
                    "region": region_name,
                    "regionConfidence": region["confidence"],
                    "regionSource": "layout",
                }
                if block_type in {
                    "formula",
                    "stamp",
                    "signature",
                    "handwriting",
                    "chart",
                    "heading",
                    "header",
                    "footer",
                }:
                    target["type"] = block_type
                continue
            if block_type in {
                "chart",
                "formula",
                "stamp",
                "signature",
                "handwriting",
                "image",
            }:
                region_block = dir_block(
                    f"p{page_number}-r{len(blocks)}",
                    block_type,
                    "",
                    region["bbox"],
                    region["confidence"],
                    source,
                    len(blocks),
                )
                region_block["meta"] = {
                    "region": region_name,
                    "regionConfidence": region["confidence"],
                    "regionSource": "layout",
                }
                blocks.append(region_block)
        blocks.sort(
            key=lambda block: (
                (block.get("bbox") or {}).get("y", float("inf")),
                (block.get("bbox") or {}).get("x", 0.0),
            )
        )
        for index, block in enumerate(blocks):
            block["order"] = index
        ink_gaps = detect_page_ink_gaps(path, blocks, page_number)
        meta: dict[str, Any] = {}
        if ink_gaps is not None:
            meta["inkGaps"] = ink_gaps
        page = dir_page(
            page_number,
            width,
            height,
            blocks,
            source,
            meta or None,
        )
        attach_machine_readable(page, machine_readable_for_page(path, text, page_number))
        return PageExtraction(normalize_ocr_text(text), overall_confidence, page, label)

    def _native_pdf_extraction(
        self, native_pages: list[dict[str, Any]], truncated: bool = False
    ) -> Extraction:
        texts: list[str] = []
        dir_pages: list[dict[str, Any]] = []
        page_trusts: list[tuple[float, list[str]]] = []
        for number, page in enumerate(native_pages, 1):
            text = page.get("text") or ""
            page_trusts.append(page_text_trust(page))
            if usable_text_layer(text):
                texts.append(normalize_ocr_text(text))
                dir_pages.append(native_pdf_dir_page(page, number))
            else:
                empty = dir_page(number, page.get("width"), page.get("height"), [], None)
                attach_machine_readable(empty, parse_mrz(text, number))
                dir_pages.append(empty)
        document = dir_document(
            "native",
            dir_pages,
            {"extractor": "pdfplumber", "textLayer": text_layer_summary(page_trusts)},
        )
        return Extraction(
            normalize_ocr_text("\n\n".join(texts).strip()),
            None,
            "pdf-text",
            self.model,
            len(native_pages),
            truncated,
            "native",
            document,
        )

    @staticmethod
    def _pdfium_text(page: Any) -> str:
        try:
            text_page = page.get_textpage()
            try:
                return text_page.get_text_bounded() or ""
            finally:
                text_page.close()
        except Exception:  # noqa: BLE001 - a broken text layer is not fatal
            return ""

    @staticmethod
    def _render_pdf_page(page: Any, dpi: int) -> Path:
        try:
            page_width, page_height = page.get_size()
            pixels = (float(page_width) * dpi / 72) * (float(page_height) * dpi / 72)
            if pixels > MAX_IMAGE_PIXELS:
                raise RequestError(
                    413,
                    f"rendered page exceeds the {MAX_IMAGE_PIXELS // 1_000_000} megapixel limit",
                )
        except RequestError:
            raise
        except Exception:  # noqa: BLE001 - size metadata is optional
            pass
        bitmap = page.render(scale=dpi / 72)
        try:
            image = bitmap.to_pil()
        finally:
            bitmap.close()
        temporary = temp_image_path()
        image.save(temporary, format="PNG")
        return temporary

    @staticmethod
    def _classify_image_region(
        bbox: dict[str, float], page_width: float, page_height: float, ratio: float
    ) -> str:
        """Geometric hint for region-level routing; never used to rewrite text."""
        width = max(1.0, float(bbox.get("width") or 0))
        height = max(1.0, float(bbox.get("height") or 0))
        aspect = width / height
        bottom = float(bbox.get("y") or 0) + height
        lower_half = bottom >= page_height * 0.55
        if aspect >= 2.0 and ratio >= 0.10:
            return "chart"
        if lower_half and ratio <= 0.15 and aspect <= 1.5:
            return "stamp"
        if lower_half and aspect >= 2.0 and ratio <= 0.08:
            return "signature"
        if ratio >= 0.2 and aspect >= 1.4:
            return "figure"
        return "image"

    @staticmethod
    def _embedded_image_regions(native: dict[str, Any]) -> list[dict[str, Any]]:
        """Image regions inside a native page that need OCR, in PDF points."""
        width = float(native.get("width") or 0)
        height = float(native.get("height") or 0)
        if width <= 0 or height <= 0:
            return []
        area = width * height
        min_ratio = env_float("OCR_HYBRID_IMAGE_MIN_RATIO", 0.04)
        special_min_ratio = env_float("OCR_HYBRID_SPECIAL_MIN_RATIO", 0.004)
        limit = env_int("OCR_HYBRID_IMAGE_MAX_REGIONS", 2)
        special_limit = env_int("OCR_HYBRID_SPECIAL_MAX_REGIONS", 2)
        lines = native.get("lines") or []
        regions: list[dict[str, Any]] = []
        special = 0
        for image in native.get("images") or []:
            bbox = image.get("bbox")
            if not isinstance(bbox, dict):
                continue
            ratio = (bbox["width"] * bbox["height"]) / area
            if ratio >= 0.8:
                continue
            kind = PaddleEngine._classify_image_region(bbox, width, height, ratio)
            is_special = kind in {"stamp", "signature"}
            floor = special_min_ratio if is_special else min_ratio
            if ratio < floor:
                continue
            if bbox["width"] < 20 or bbox["height"] < 20:
                continue
            if sum(1 for line in lines if _line_inside_bbox(line, bbox)) >= 3:
                continue
            if is_special and special >= special_limit:
                continue
            regions.append({**bbox, "type": kind})
            if is_special:
                special += 1
            if len(regions) >= limit + special_limit:
                break
        return regions

    def _ocr_embedded_images(
        self,
        pdf_page: Any,
        native: dict[str, Any],
        number: int,
        dpi: int,
        lang: str | None = None,
        model: str | None = None,
    ) -> tuple[list[dict[str, Any]], str, list[float]]:
        """OCR raster regions embedded in a page that already has native text."""
        regions = self._embedded_image_regions(native)
        if not regions:
            return [], "", []
        rendered = self._render_pdf_page(pdf_page, dpi)
        try:
            try:
                from PIL import Image
            except Exception:  # noqa: BLE001 - Pillow is required for crops
                return [], "", []
            with Image.open(rendered) as opened:
                image = opened.convert("RGB")
            page_width = float(native.get("width") or image.width)
            page_height = float(native.get("height") or image.height)
            scale_x = image.width / page_width if page_width else 1.0
            scale_y = image.height / page_height if page_height else 1.0
            blocks: list[dict[str, Any]] = []
            texts: list[str] = []
            scores: list[float] = []
            for index, region in enumerate(regions):
                left = max(0, int(region["x"] * scale_x))
                top = max(0, int(region["y"] * scale_y))
                right = min(image.width, int((region["x"] + region["width"]) * scale_x + 0.5))
                bottom = min(image.height, int((region["y"] + region["height"]) * scale_y + 0.5))
                if right - left < 20 or bottom - top < 20:
                    continue
                temporary = temp_image_path()
                try:
                    image.crop((left, top, right, bottom)).save(temporary, format="PNG")
                    page_data = self._page_from_image(temporary, "auto", number, lang, model)
                finally:
                    temporary.unlink(missing_ok=True)
                if page_data.confidence is not None:
                    scores.append(page_data.confidence)
                if page_data.text.strip():
                    texts.append(page_data.text)
                for block in page_data.page.get("blocks") or []:
                    bbox = block.get("bbox")
                    if isinstance(bbox, dict):
                        block["bbox"] = {
                            "x": region["x"] + bbox["x"] / scale_x,
                            "y": region["y"] + bbox["y"] / scale_y,
                            "width": bbox["width"] / scale_x,
                            "height": bbox["height"] / scale_y,
                        }
                    block["id"] = f"p{number}-img{index}-{block.get('id', 'b0')}"
                    region_type = str(region.get("type") or "image")
                    block["meta"] = {
                        **(block.get("meta") or {}),
                        "hybrid": f"embedded_{region_type}",
                        "regionType": region_type,
                        # Bboxes above were mapped back to PDF points.
                        "bboxDpi": 72,
                    }
                    block["order"] = len(blocks)
                    blocks.append(block)
            return blocks, "\n\n".join(text for text in texts if text.strip()), scores
        except Exception as cause:  # noqa: BLE001 - hybrid pages are best-effort
            print(
                f"embedded image OCR failed on page {number} ({cause}); keeping native text",
                file=sys.stderr,
            )
            return [], "", []
        finally:
            rendered.unlink(missing_ok=True)

    def _crosscheck_text_layer(
        self,
        document: Any,
        native_pages: list[dict[str, Any]],
        positions: list[int],
        dpi: int,
        lang: str | None,
        model: str | None,
        started: float,
        budget: float,
    ) -> dict[str, Any]:
        """Spot-check trusted native pages against a fresh OCR render.

        The native text is never replaced; a mismatch only becomes quality
        evidence so reviewers know the embedded layer may be stale or wrong.
        """
        threshold = text_crosscheck_min_similarity()
        selected = select_crosscheck_pages(
            positions,
            text_crosscheck_rate(),
            text_crosscheck_max_pages(),
            crosscheck_priority_pages(native_pages, positions),
        )
        samples: list[dict[str, Any]] = []
        failed = 0
        for position in selected:
            if time.perf_counter() - started >= budget:
                break
            number = position + 1
            try:
                temporary = self._render_pdf_page(document[position], dpi)
                try:
                    observed = self._page_from_image(temporary, "auto", number, lang, model)
                finally:
                    temporary.unlink(missing_ok=True)
                samples.append(
                    crosscheck_record(
                        number,
                        str(native_pages[position].get("text") or ""),
                        observed.text,
                        threshold,
                    )
                )
            except Exception as cause:  # noqa: BLE001 - sampling must not fail extraction
                failed += 1
                print(
                    f"text-layer crosscheck failed on page {number} ({cause})",
                    file=sys.stderr,
                )
        return crosscheck_summary(samples, threshold, len(selected), failed)

    @staticmethod
    def _crosscheck_skipped(reason: str) -> dict[str, Any]:
        """Report a skipped crosscheck instead of silently passing the layer."""
        return {
            "sampled": 0,
            "planned": 0,
            "failed": 0,
            "matched": 0,
            "mismatched": 0,
            "numericMismatched": 0,
            "threshold": text_crosscheck_min_similarity(),
            "minSimilarity": None,
            "pages": [],
            "budgetExceeded": False,
            "skipped": reason[:200],
        }

    def _crosscheck_native(
        self,
        path: Path,
        native_pages: list[dict[str, Any]],
        dpi: int,
        lang: str | None,
        model: str | None,
    ) -> dict[str, Any] | None:
        """Crosscheck the all-native fast path with a bounded render sample."""
        positions = [
            index for index, page in enumerate(native_pages) if trustworthy_native_page(page)
        ]
        if not select_crosscheck_pages(
            positions,
            text_crosscheck_rate(),
            text_crosscheck_max_pages(),
            crosscheck_priority_pages(native_pages, positions),
        ):
            return None
        try:
            import pypdfium2 as pdfium
        except Exception as cause:  # noqa: BLE001 - crosscheck is best-effort
            print(
                f"pypdfium2 unavailable for text-layer crosscheck ({cause})",
                file=sys.stderr,
            )
            return self._crosscheck_skipped(f"pypdfium2 unavailable: {cause}")
        try:
            document = pdfium.PdfDocument(str(path))
        except Exception as cause:  # noqa: BLE001 - crosscheck is best-effort
            print(f"cannot open PDF for text-layer crosscheck ({cause})", file=sys.stderr)
            return self._crosscheck_skipped(f"cannot open PDF: {cause}")
        started = time.perf_counter()
        budget = env_int("OCR_PDF_TIME_BUDGET_SECONDS", DEFAULT_PDF_TIME_BUDGET_SECONDS)
        try:
            return self._crosscheck_text_layer(
                document, native_pages, positions, dpi, lang, model, started, budget
            )
        finally:
            try:
                document.close()
            except Exception:  # noqa: BLE001 - already closed
                pass

    def _extract_pdf(
        self,
        path: Path,
        dpi_override: int | None = None,
        mode: str = "auto",
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> Extraction:
        native_pages = read_pdfplumber_pages(
            path, env_int("OCR_MAX_PDF_TOTAL_PAGES", DEFAULT_MAX_PDF_TOTAL_PAGES)
        )
        if lang is None and native_pages:
            # Only trusted text may pick the language profile; a placeholder or
            # corrupted layer must not decide which model renders scanned pages.
            trusted_text = "\n".join(
                str(page.get("text") or "")
                for page in native_pages
                if trustworthy_native_page(page)
            )
            if trusted_text:
                lang = detect_language(trusted_text)
        force_visual = mode in {"table", "structure"}
        if (
            native_pages
            and not force_visual
            and all(trustworthy_native_page(page) for page in native_pages)
        ):
            extraction = self._native_pdf_extraction(native_pages)
            crosscheck = self._crosscheck_native(
                path,
                native_pages,
                dpi_override or env_int("OCR_PDF_DPI", DEFAULT_PDF_DPI),
                lang,
                model,
            )
            if crosscheck is not None and extraction.document:
                text_layer = extraction.document.get("meta", {}).get("textLayer")
                if isinstance(text_layer, dict):
                    text_layer["crosscheck"] = crosscheck
            return extraction

        try:
            import pypdfium2 as pdfium
        except Exception as cause:  # noqa: BLE001 - fall back to Paddle's PDF reader
            print(f"pypdfium2 unavailable ({cause}); OCR-ing the PDF directly", file=sys.stderr)
            if native_pages:
                return self._native_pdf_extraction(native_pages)
            return self._ocr_pdf(path, lang, model)

        try:
            document = pdfium.PdfDocument(str(path))
        except Exception as cause:  # noqa: BLE001 - encrypted or malformed PDF
            print(f"cannot open PDF text layer ({cause}); OCR-ing directly", file=sys.stderr)
            if native_pages:
                return self._native_pdf_extraction(native_pages)
            return self._ocr_pdf(path, lang, model)

        dpi = dpi_override or env_int("OCR_PDF_DPI", DEFAULT_PDF_DPI)
        max_ocr_pages = env_int("OCR_MAX_PDF_PAGES", DEFAULT_MAX_PDF_PAGES)
        max_total_pages = env_int("OCR_MAX_PDF_TOTAL_PAGES", DEFAULT_MAX_PDF_TOTAL_PAGES)
        budget = env_int("OCR_PDF_TIME_BUDGET_SECONDS", DEFAULT_PDF_TIME_BUDGET_SECONDS)
        started = time.perf_counter()
        total = len(document)
        pages_text: list[str] = []
        dir_pages: list[dict[str, Any]] = []
        scores: list[float] = []
        used_native = False
        used_ocr = False
        ocr_pages = 0
        skipped_pages = 0
        native_kept: list[int] = []
        crosscheck: dict[str, Any] | None = None
        preprocess_pages: list[dict[str, Any]] = []
        truncated = total > max_total_pages
        try:
            for index in range(min(total, max_total_pages)):
                number = index + 1
                if index > 0 and time.perf_counter() - started >= budget:
                    truncated = True
                    break
                try:
                    page = document[index]
                    native = (
                        native_pages[index]
                        if native_pages and index < len(native_pages)
                        else None
                    )
                    native_text = (native.get("text") or "") if native else ""
                    native_tables = (native.get("tables") or []) if native else []
                    # A trusted native layer is never replaced by a visual
                    # structure/table rerun; only untrusted (scanned) pages are
                    # OCR'd in those modes. A forced visual pass on a trusted
                    # page may still contribute auxiliary table blocks.
                    if native and trustworthy_native_page(native):
                        dir_value = native_pdf_dir_page(native, number)
                        page_texts = [normalize_ocr_text(native_text)]
                        if (
                            force_visual
                            and not native_tables
                            and time.perf_counter() - started < budget
                        ):
                            try:
                                temporary = self._render_pdf_page(page, dpi)
                                try:
                                    page_data = self._page_from_image(
                                        temporary, mode, number, lang, model, channel
                                    )
                                finally:
                                    temporary.unlink(missing_ok=True)
                                scale = 72.0 / dpi
                                merged = 0
                                for block in page_data.page.get("blocks") or []:
                                    if block.get("type") != "table":
                                        continue
                                    bbox = block.get("bbox")
                                    if isinstance(bbox, dict):
                                        block["bbox"] = {
                                            key: float(value) * scale
                                            for key, value in bbox.items()
                                        }
                                    for cell in block.get("cells") or []:
                                        cell_bbox = cell.get("bbox")
                                        if isinstance(cell_bbox, dict):
                                            cell["bbox"] = {
                                                key: float(value) * scale
                                                for key, value in cell_bbox.items()
                                            }
                                    block["id"] = (
                                        f"p{number}-aux{merged}-{block.get('id', 't0')}"
                                    )
                                    block["meta"] = {
                                        **(block.get("meta") or {}),
                                        "hybrid": "auxiliary_table",
                                        "bboxDpi": 72,
                                    }
                                    block["order"] = len(dir_value["blocks"])
                                    dir_value["blocks"].append(block)
                                    merged += 1
                                if merged:
                                    used_ocr = True
                                    if page_data.confidence is not None:
                                        scores.append(page_data.confidence)
                            except Exception as cause:  # noqa: BLE001 - keep native text
                                print(
                                    f"auxiliary table pass failed on page {number} ({cause}); "
                                    "keeping native text",
                                    file=sys.stderr,
                                )
                        elif not force_visual and time.perf_counter() - started < budget:
                            extra_blocks, extra_text, extra_scores = (
                                self._ocr_embedded_images(page, native, number, dpi, lang, model)
                            )
                            if extra_blocks:
                                dir_value["blocks"].extend(extra_blocks)
                                dir_value["blocks"].sort(
                                    key=lambda block: (
                                        (block.get("bbox") or {}).get("y", float("inf")),
                                        (block.get("bbox") or {}).get("x", 0.0),
                                    )
                                )
                                for order, block in enumerate(dir_value["blocks"]):
                                    block["order"] = order
                                if extra_text:
                                    page_texts.append(extra_text)
                                scores.extend(extra_scores)
                                used_ocr = True
                        pages_text.append("\n\n".join(page_texts))
                        dir_pages.append(dir_value)
                        used_native = True
                        native_kept.append(index)
                        continue

                    if native is None and not force_visual:
                        page_text = self._pdfium_text(page)
                        if usable_text_layer(page_text):
                            pages_text.append(normalize_ocr_text(page_text))
                            source = dir_source("pypdfium2", None, "native")
                            blocks = [
                                dir_block(
                                    f"p{number}-b0",
                                    "paragraph",
                                    page_text,
                                    None,
                                    None,
                                    source,
                                    0,
                                )
                            ]
                            dir_pages.append(
                                dir_page(number, None, None, blocks, source)
                            )
                            used_native = True
                            continue

                    # Digital pages above are free; only rendered pages are
                    # capped, so long text PDFs are not cut at 100 pages.
                    if ocr_pages >= max_ocr_pages:
                        truncated = True
                        skipped_pages += 1
                        dir_pages.append(
                            dir_page(
                                number,
                                float(getattr(page, "get_width", lambda: 0)()),
                                float(getattr(page, "get_height", lambda: 0)()),
                                [],
                                None,
                            )
                        )
                        continue
                    ocr_pages += 1

                    temporary = self._render_pdf_page(page, dpi)
                    try:
                        page_mode = mode if force_visual or mode == "enhanced" else "auto"
                        page_data = self._page_from_image(
                            temporary, page_mode, number, lang, model, channel
                        )
                        pages_text.append(page_data.text)
                        dir_pages.append(page_data.page)
                        if page_data.preprocess:
                            preprocess_pages.append(page_data.preprocess)
                        if page_data.confidence is not None:
                            scores.append(page_data.confidence)
                        used_ocr = True
                    finally:
                        temporary.unlink(missing_ok=True)
                except Exception as cause:  # noqa: BLE001 - one bad page should not lose the rest
                    print(f"page {index + 1}/{total} failed ({cause}); skipping", file=sys.stderr)
                    if not dir_pages or dir_pages[-1].get("number") != number:
                        dir_pages.append(dir_page(number, None, None, [], None))
            if native_kept:
                crosscheck = self._crosscheck_text_layer(
                    document, native_pages, native_kept, dpi, lang, model, started, budget
                )
        finally:
            try:
                document.close()
            except Exception:  # noqa: BLE001 - already closed
                pass

        text = "\n\n".join(page for page in pages_text if page.strip()).strip()
        if not text:
            return self._ocr_pdf(path, lang, model)
        # Truncation is metadata, never in-band text: downstream search, field
        # matching and CER evaluation must not read the progress note as content.
        confidence = sum(scores) / len(scores) if scores else None
        base_model = self._model_for(lang, model)
        if used_native and used_ocr:
            route = "hybrid"
            label = f"{base_model}+pdf-text"
        elif used_native:
            route = "native"
            label = "pdf-text"
        else:
            route = "visual"
            label = base_model
        meta: dict[str, Any] = {
            "extractor": "pdfplumber" if native_pages else "pdfium",
            "mode": mode,
        }
        if truncated:
            meta["truncated"] = True
            meta["processedPages"] = len(pages_text)
            meta["totalPages"] = total
            if skipped_pages:
                meta["skippedScanPages"] = skipped_pages
        if channel != "rgb":
            meta["channel"] = channel
        if native_pages:
            text_layer = text_layer_summary([page_text_trust(page) for page in native_pages])
            if crosscheck is not None:
                text_layer["crosscheck"] = crosscheck
            meta["textLayer"] = text_layer
        if lang:
            meta["lang"] = lang
        if preprocess_pages:
            meta["preprocess"] = {"pages": preprocess_pages}
        document_ir = dir_document(route, dir_pages, meta)
        return Extraction(
            text,
            confidence,
            label,
            base_model,
            len(pages_text),
            truncated,
            route,
            document_ir,
        )

    def _ocr_pdf(
        self, path: Path, lang: str | None = None, model: str | None = None
    ) -> Extraction:
        ocr = self._profile_instance(model, lang, False)
        base_model = self._model_for(lang, model)
        max_pages = env_int("OCR_MAX_PDF_TOTAL_PAGES", DEFAULT_MAX_PDF_TOTAL_PAGES)
        budget = env_int("OCR_PDF_TIME_BUDGET_SECONDS", DEFAULT_PDF_TIME_BUDGET_SECONDS)
        started = time.perf_counter()
        texts: list[str] = []
        dir_pages: list[dict[str, Any]] = []
        scores: list[float] = []
        truncated = False
        for index, result in enumerate(ocr.predict(input=str(path)), 1):
            if index > max_pages or time.perf_counter() - started >= budget:
                truncated = True
                break
            text, confidence, blocks, width, height = paddle_page_blocks(
                result, index, "visual", base_model, base_model
            )
            texts.append(text)
            dir_pages.append(
                dir_page(index, width, height, blocks, dir_source(base_model, base_model, "visual"))
            )
            if confidence is not None:
                scores.append(confidence)
        text = "\n\n".join(page for page in texts if page.strip()).strip()
        confidence = sum(scores) / len(scores) if scores else None
        meta = {"truncated": True, "processedPages": len(texts)} if truncated else None
        return Extraction(
            normalize_ocr_text(text),
            confidence,
            base_model,
            base_model,
            len(texts),
            truncated,
            "visual",
            dir_document("visual", dir_pages, meta),
        )


def _clean_rows(
    rows: Any,
    max_rows: int = 5000,
    max_columns: int = 200,
    flags: dict[str, bool] | None = None,
) -> list[list[str]]:
    cleaned: list[list[str]] = []
    for row in rows:
        if isinstance(row, (list, tuple)):
            if flags is not None and len(row) > max_columns:
                flags["truncated"] = True
            values = ["" if cell is None else str(cell).strip() for cell in row[:max_columns]]
        else:
            values = ["" if row is None else str(row).strip()]
        if any(values):
            cleaned.append(values)
        if len(cleaned) >= max_rows:
            if flags is not None:
                flags["truncated"] = True
            break
    return cleaned


def _local_name(tag: Any) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


MAX_XML_BYTES = 32 * 1024 * 1024


def _safe_xml(data: bytes) -> Any:
    """Reject DTD/entity declarations before parsing an Office part."""
    if len(data) > MAX_XML_BYTES:
        raise ValueError("XML part exceeds the parse limit")
    head = data[:8192].upper()
    if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
        raise ValueError("XML DTD/entity declarations are not allowed")
    return ET.fromstring(data)


def _docx_element_text(element: Any) -> tuple[str, dict[str, int]]:
    """Word text with tracked deletions excluded and insertions kept."""
    parts: list[str] = []
    stats = {"deleted": 0, "inserted": 0, "changes": 0}

    def walk(node: Any, deleted: bool, inserted: bool) -> None:
        for child in node:
            name = _local_name(child.tag)
            # Word repeats drawings inside mc:Fallback; reading both duplicates
            # text boxes and embedded images.
            if name == "Fallback" and _local_name(node.tag) == "AlternateContent":
                continue
            next_deleted = deleted or name in {"del", "moveFrom"}
            next_inserted = inserted or name in {"ins", "moveTo"}
            if name in {"ins", "del", "moveFrom", "moveTo"}:
                stats["changes"] += 1
            if name in {"t", "delText"}:
                value = child.text or ""
                if next_deleted:
                    stats["deleted"] += len(value)
                else:
                    parts.append(value)
                    if next_inserted:
                        stats["inserted"] += len(value)
                continue
            if name == "tab":
                if not next_deleted:
                    parts.append("\t")
                continue
            if name in {"br", "cr"}:
                if not next_deleted:
                    parts.append("\n")
                continue
            if name == "noBreakHyphen":
                if not next_deleted:
                    parts.append("-")
                continue
            walk(child, next_deleted, next_inserted)

    walk(element, False, False)
    return "".join(parts), stats


def _docx_paragraph_texts(element: Any) -> list[str]:
    found: list[str] = []
    for paragraph in element.iter():
        if _local_name(paragraph.tag) != "p":
            continue
        text = _docx_element_text(paragraph)[0].strip()
        if text:
            found.append(text)
    return found


def _docx_table_rows(table: Any, flags: dict[str, bool]) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                cells.append("")
                continue
            seen.add(key)
            cells.append(
                "\n".join(
                    _docx_element_text(paragraph._p)[0].strip()
                    for paragraph in cell.paragraphs
                ).strip()
            )
        rows.append(cells)
    return _clean_rows(rows, flags=flags)


def _chart_text(data: bytes) -> dict[str, Any]:
    root = _safe_xml(data)
    titles: list[str] = []
    values: list[str] = []
    for node in root.iter():
        local = _local_name(node.tag)
        value = (node.text or "").strip()
        if not value:
            continue
        if local == "t":
            titles.append(value)
        elif local == "v":
            values.append(value)
    return {
        "text": " | ".join(dict.fromkeys([*titles[:3], *values[:60]]))[:4000],
        "values": len(values),
    }


def _office_zip_parts(path: Path) -> dict[str, list[dict[str, Any]]]:
    """Charts, media, comments and OLE payloads for any OOXML package."""
    result: dict[str, list[dict[str, Any]]] = {
        "charts": [],
        "media": [],
        "embeddings": [],
        "comments": [],
    }
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if re.fullmatch(r"(?:word|xl|ppt)/charts?/chart\d+\.xml", name):
                chart = _chart_text(archive.read(name))
                result["charts"].append(
                    {"part": name, "text": chart["text"], "values": chart["values"]}
                )
            elif re.match(r"(?:word|xl|ppt)/media/", name):
                result["media"].append({"name": name, "bytes": archive.getinfo(name).file_size})
            elif re.match(r"(?:word|xl|ppt)/embeddings/", name):
                result["embeddings"].append(
                    {"name": name, "bytes": archive.getinfo(name).file_size}
                )
            elif re.fullmatch(r"(?:xl|ppt)/comments\d*\.xml", name):
                root = _safe_xml(archive.read(name))
                text = " ".join(
                    (node.text or "").strip()
                    for node in root.iter()
                    if _local_name(node.tag) == "t" and (node.text or "").strip()
                ).strip()
                if text:
                    result["comments"].append({"part": name, "text": text[:4000]})
    return result


def _docx_side_parts(path: Path) -> dict[str, Any]:
    """Footnotes, endnotes, comments, charts and embedded objects from the zip."""
    result: dict[str, Any] = {
        "footnotes": [],
        "endnotes": [],
        "comments": [],
        "charts": [],
        "media": [],
        "embeddings": [],
    }
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        for part, key in (
            ("word/footnotes.xml", "footnote"),
            ("word/endnotes.xml", "endnote"),
            ("word/comments.xml", "comment"),
        ):
            if part not in names:
                continue
            root = _safe_xml(archive.read(part))
            for item in root:
                kind = _local_name(item.tag)
                if kind not in {key, "comment"}:
                    continue
                if kind == "footnote" and item.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
                ) in {"separator", "continuationSeparator", "continuationNotice"}:
                    continue
                text = " ".join(_docx_paragraph_texts(item)).strip()
                if not text:
                    continue
                entry: dict[str, Any] = {"text": text}
                if kind == "comment":
                    entry["author"] = item.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                    )
                    entry["date"] = item.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"
                    )
                result[f"{kind}s"].append(entry)
        for name in names:
            if re.fullmatch(r"word/charts/chart\d+\.xml", name):
                root = _safe_xml(archive.read(name))
                titles: list[str] = []
                values: list[str] = []
                for node in root.iter():
                    local = _local_name(node.tag)
                    value = (node.text or "").strip()
                    if not value:
                        continue
                    if local == "t":
                        titles.append(value)
                    elif local == "v":
                        values.append(value)
                text = " | ".join(dict.fromkeys([*titles[:3], *values[:60]]))[:4000]
                result["charts"].append({"part": name, "text": text, "values": len(values)})
            elif name.startswith("word/media/"):
                result["media"].append({"name": name, "bytes": archive.getinfo(name).file_size})
            elif name.startswith("word/embeddings/"):
                result["embeddings"].append(
                    {"name": name, "bytes": archive.getinfo(name).file_size}
                )
    return result


def extract_docx(path: Path) -> NativeText | None:
    try:
        from docx import Document
    except Exception:  # noqa: BLE001 - optional native extractor
        return None
    try:
        document = Document(str(path))
        source = dir_source("python-docx", None, "native")
        blocks: list[dict[str, Any]] = []
        texts: list[str] = []
        flags = {"truncated": False}
        stats = {"deleted": 0, "inserted": 0, "changes": 0}

        def add(block_type: str, text: str, meta: dict[str, Any] | None = None) -> None:
            text = text.strip()
            if not text:
                return
            block = dir_block(
                f"p1-b{len(blocks)}",
                block_type,
                text,
                None,
                None,
                source,
                len(blocks),
            )
            if meta:
                block["meta"] = meta
            blocks.append(block)
            texts.append(text)

        for paragraph in document.paragraphs:
            text, paragraph_stats = _docx_element_text(paragraph._p)
            for key in stats:
                stats[key] += paragraph_stats[key]
            style = (getattr(paragraph.style, "name", "") or "").lower()
            block_type = "heading" if style.startswith(("heading", "标题")) else "paragraph"
            add(block_type, text)

        for table in document.tables:
            rows = _docx_table_rows(table, flags)
            if not rows:
                continue
            table_text = "\n".join(" | ".join(row) for row in rows)
            add("table", table_text)
            blocks[-1]["rows"] = rows

        header_count = 0
        footer_count = 0
        seen_layout: set[str] = set()
        for section in document.sections:
            for kind, container in (
                ("header", section.header),
                ("footer", section.footer),
                ("header", section.first_page_header),
                ("footer", section.first_page_footer),
                ("header", section.even_page_header),
                ("footer", section.even_page_footer),
            ):
                if container.is_linked_to_previous:
                    continue
                for paragraph in container.paragraphs:
                    text, part_stats = _docx_element_text(paragraph._p)
                    for key in stats:
                        stats[key] += part_stats[key]
                    text = text.strip()
                    if not text or text in seen_layout:
                        continue
                    seen_layout.add(text)
                    add(kind, text, {"region": kind})
                    if kind == "header":
                        header_count += 1
                    else:
                        footer_count += 1
                for table in container.tables:
                    rows = _docx_table_rows(table, flags)
                    table_text = "\n".join(" | ".join(row) for row in rows).strip()
                    if not table_text or table_text in seen_layout:
                        continue
                    seen_layout.add(table_text)
                    add(kind, table_text, {"region": kind})

        side = _docx_side_parts(path)
        for entry in side["footnotes"]:
            add("paragraph", entry["text"], {"region": "footnote"})
        for entry in side["endnotes"]:
            add("paragraph", entry["text"], {"region": "endnote"})
        for entry in side["comments"]:
            add(
                "paragraph",
                entry["text"],
                {"region": "comment", "author": entry.get("author"), "date": entry.get("date")},
            )
        for chart in side["charts"]:
            add("chart", chart["text"], {"chart": chart["part"], "values": chart["values"]})
        for image in side["media"]:
            add("image", "", {"embedded": "image", "media": image["name"], "bytes": image["bytes"]})
        for embedded in side["embeddings"]:
            add(
                "image",
                "",
                {"embedded": "ole", "media": embedded["name"], "bytes": embedded["bytes"]},
            )

        textboxes = sum(
            1 for node in document.element.body.iter() if _local_name(node.tag) == "txbxContent"
        )
        office = {
            "deletedText": stats["deleted"],
            "insertedText": stats["inserted"],
            "revisions": stats["changes"],
            "textBoxes": textboxes,
            "headers": header_count,
            "footers": footer_count,
            "footnotes": len(side["footnotes"]),
            "endnotes": len(side["endnotes"]),
            "comments": len(side["comments"]),
            "charts": len(side["charts"]),
            "images": len(side["media"]),
            "embeddedObjects": len(side["embeddings"]),
            "hiddenSlides": 0,
            "groupedShapes": 0,
            "notes": 0,
            "objects": len(side["charts"]) + len(side["media"]) + len(side["embeddings"]),
            "unrendered": len(side["media"]) + len(side["embeddings"]),
        }
        page = dir_page(1, None, None, blocks, source)
        attach_machine_readable(page, parse_mrz("\n".join(texts), 1))
        return NativeText(
            "\n".join(texts), [page], flags["truncated"], {"office": office}
        )
    except Exception as cause:  # noqa: BLE001 - fall back to Docling
        print(f"python-docx failed ({cause}); falling back to Docling", file=sys.stderr)
        return None


def extract_xlsx(path: Path) -> NativeText | None:
    try:
        from openpyxl import load_workbook
    except Exception:  # noqa: BLE001 - optional native extractor
        return None
    workbook = None
    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        source = dir_source("openpyxl", None, "native")
        pages: list[dict[str, Any]] = []
        texts: list[str] = []
        flags = {"truncated": False}
        hidden_sheets = 0
        max_pages = env_int("OCR_MAX_OFFICE_PAGES", DEFAULT_MAX_OFFICE_PAGES)
        for number, sheet in enumerate(workbook.worksheets, 1):
            if number > max_pages:
                flags["truncated"] = True
                break
            rows = _clean_rows(sheet.iter_rows(values_only=True), flags=flags)
            hidden = str(getattr(sheet, "sheet_state", "visible")) != "visible"
            if hidden:
                hidden_sheets += 1
            if not rows and not hidden:
                continue
            heading = dir_block(
                f"p{number}-b0",
                "heading",
                str(sheet.title),
                None,
                None,
                source,
                0,
            )
            if hidden:
                heading["meta"] = {"hidden": True}
            blocks = [heading]
            table_text = "\n".join(" | ".join(row) for row in rows)
            if table_text:
                blocks.append(
                    dir_block(
                        f"p{number}-t1",
                        "table",
                        table_text,
                        None,
                        None,
                        source,
                        1,
                        rows=rows,
                    )
                )
            page = dir_page(number, None, None, blocks, source)
            attach_machine_readable(page, parse_mrz(table_text, number))
            pages.append(page)
            if table_text:
                texts.append(f"## {sheet.title}\n{table_text}")
            else:
                texts.append(f"## {sheet.title}")

        side = _office_zip_parts(path)
        objects = side["charts"] + side["media"] + side["embeddings"]
        if side["comments"] or objects:
            if not pages:
                pages.append(dir_page(1, None, None, [], source))
            carrier = pages[0]
            for comment in side["comments"]:
                block = dir_block(
                    f"p{carrier['number']}-c{len(carrier['blocks'])}",
                    "paragraph",
                    comment["text"],
                    None,
                    None,
                    source,
                    len(carrier["blocks"]),
                )
                block["meta"] = {"region": "comment", "part": comment["part"]}
                carrier["blocks"].append(block)
                texts.append(comment["text"])
            for chart in side["charts"]:
                block = dir_block(
                    f"p{carrier['number']}-c{len(carrier['blocks'])}",
                    "chart",
                    chart["text"],
                    None,
                    None,
                    source,
                    len(carrier["blocks"]),
                )
                block["meta"] = {"chart": chart["part"], "values": chart["values"]}
                carrier["blocks"].append(block)
                if chart["text"]:
                    texts.append(chart["text"])
            for item in side["media"]:
                block = dir_block(
                    f"p{carrier['number']}-i{len(carrier['blocks'])}",
                    "image",
                    "",
                    None,
                    None,
                    source,
                    len(carrier["blocks"]),
                )
                block["meta"] = {
                    "embedded": "image",
                    "media": item["name"],
                    "bytes": item["bytes"],
                }
                carrier["blocks"].append(block)
            for item in side["embeddings"]:
                block = dir_block(
                    f"p{carrier['number']}-o{len(carrier['blocks'])}",
                    "image",
                    "",
                    None,
                    None,
                    source,
                    len(carrier["blocks"]),
                )
                block["meta"] = {
                    "embedded": "ole",
                    "media": item["name"],
                    "bytes": item["bytes"],
                }
                carrier["blocks"].append(block)

        office = {
            "deletedText": 0,
            "insertedText": 0,
            "revisions": 0,
            "textBoxes": 0,
            "headers": 0,
            "footers": 0,
            "footnotes": 0,
            "endnotes": 0,
            "comments": len(side["comments"]),
            "charts": len(side["charts"]),
            "images": len(side["media"]),
            "embeddedObjects": len(side["embeddings"]),
            "hiddenSlides": hidden_sheets,
            "groupedShapes": 0,
            "notes": 0,
            "objects": len(objects),
            "unrendered": len(side["media"]) + len(side["embeddings"]),
        }
        return NativeText(
            "\n\n".join(texts), pages, flags["truncated"], {"office": office}
        )
    except Exception as cause:  # noqa: BLE001 - fall back to Docling
        print(f"openpyxl failed ({cause}); falling back to Docling", file=sys.stderr)
        return None
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:  # noqa: BLE001 - already closed
                pass


def extract_pptx(path: Path) -> NativeText | None:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except Exception:  # noqa: BLE001 - optional native extractor
        return None
    try:
        presentation = Presentation(str(path))
        ole_type = getattr(MSO_SHAPE_TYPE, "EMBEDDED_OLE_OBJECT", None)
        source = dir_source("python-pptx", None, "native")
        pages: list[dict[str, Any]] = []
        texts: list[str] = []
        flags = {"truncated": False}
        grouped_shapes = 0
        picture_count = 0
        ole_count = 0
        notes_count = 0
        chart_count = 0
        hidden_slides = 0
        max_pages = env_int("OCR_MAX_OFFICE_PAGES", DEFAULT_MAX_OFFICE_PAGES)

        def add_shapes(
            number: int,
            shapes: Any,
            blocks: list[dict[str, Any]],
            slide_texts: list[str],
            grouped: bool = False,
        ) -> None:
            nonlocal grouped_shapes, picture_count, ole_count, chart_count
            for shape in shapes:
                if grouped:
                    grouped_shapes += 1
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    add_shapes(number, shape.shapes, blocks, slide_texts, True)
                    continue
                if getattr(shape, "has_text_frame", False):
                    text = str(shape.text_frame.text).strip()
                    if text:
                        blocks.append(
                            dir_block(
                                f"p{number}-b{len(blocks)}",
                                "paragraph",
                                text,
                                None,
                                None,
                                source,
                                len(blocks),
                            )
                        )
                        slide_texts.append(text)
                if getattr(shape, "has_table", False):
                    rows = _clean_rows(
                        [[cell.text for cell in row.cells] for row in shape.table.rows],
                        flags=flags,
                    )
                    if rows:
                        table_text = "\n".join(" | ".join(row) for row in rows)
                        blocks.append(
                            dir_block(
                                f"p{number}-t{len(blocks)}",
                                "table",
                                table_text,
                                None,
                                None,
                                source,
                                len(blocks),
                                rows=rows,
                            )
                        )
                        slide_texts.append(table_text)
                if getattr(shape, "has_chart", False):
                    chart = shape.chart
                    pieces: list[str] = []
                    try:
                        if chart.has_title:
                            pieces.append(chart.chart_title.text_frame.text.strip())
                    except Exception:  # noqa: BLE001 - title is optional
                        pass
                    try:
                        for plot in chart.plots:
                            for series in plot.series:
                                if series.name:
                                    pieces.append(str(series.name))
                                for value in list(series.values)[:40]:
                                    if value is not None:
                                        pieces.append(str(value))
                            for category in list(plot.categories)[:40]:
                                if category is not None:
                                    pieces.append(str(category))
                    except Exception:  # noqa: BLE001 - cached values are optional
                        pass
                    chart_text = " | ".join(dict.fromkeys(item for item in pieces if item))
                    block = dir_block(
                        f"p{number}-ch{len(blocks)}",
                        "chart",
                        chart_text[:4000],
                        None,
                        None,
                        source,
                        len(blocks),
                    )
                    block["meta"] = {"chart": True}
                    blocks.append(block)
                    chart_count += 1
                    if chart_text:
                        slide_texts.append(chart_text)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                    block = dir_block(
                        f"p{number}-i{len(blocks)}",
                        "image",
                        "",
                        None,
                        None,
                        source,
                        len(blocks),
                    )
                    block["meta"] = {
                        "embedded": "image",
                        "shape": getattr(shape, "name", None),
                    }
                    blocks.append(block)
                elif ole_type is not None and shape.shape_type == ole_type:
                    ole_count += 1
                    block = dir_block(
                        f"p{number}-o{len(blocks)}",
                        "image",
                        "",
                        None,
                        None,
                        source,
                        len(blocks),
                    )
                    block["meta"] = {
                        "embedded": "ole",
                        "shape": getattr(shape, "name", None),
                    }
                    blocks.append(block)

        for number, slide in enumerate(presentation.slides, 1):
            if number > max_pages:
                flags["truncated"] = True
                break
            blocks: list[dict[str, Any]] = []
            slide_texts: list[str] = []
            add_shapes(number, slide.shapes, blocks, slide_texts)
            hidden = str(slide._element.get("show")) == "0"
            if hidden:
                hidden_slides += 1
            if getattr(slide, "has_notes_slide", False):
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    notes_count += 1
                    block = dir_block(
                        f"p{number}-n{len(blocks)}",
                        "paragraph",
                        notes,
                        None,
                        None,
                        source,
                        len(blocks),
                    )
                    block["meta"] = {"region": "notes", "hidden": hidden}
                    blocks.append(block)
                    slide_texts.append(notes)
            if not blocks:
                continue
            page = dir_page(number, None, None, blocks, source)
            if hidden:
                page["meta"] = {"hidden": True}
            attach_machine_readable(page, parse_mrz("\n".join(slide_texts), number))
            pages.append(page)
            texts.append("\n".join(slide_texts))

        side = _office_zip_parts(path)
        if side["comments"]:
            if not pages:
                pages.append(dir_page(1, None, None, [], source))
            carrier = pages[0]
            for comment in side["comments"]:
                block = dir_block(
                    f"p{carrier['number']}-c{len(carrier['blocks'])}",
                    "paragraph",
                    comment["text"],
                    None,
                    None,
                    source,
                    len(carrier["blocks"]),
                )
                block["meta"] = {"region": "comment", "part": comment["part"]}
                carrier["blocks"].append(block)
                texts.append(comment["text"])

        images = max(picture_count, len(side["media"]))
        embedded = max(ole_count, len(side["embeddings"]))
        office = {
            "deletedText": 0,
            "insertedText": 0,
            "revisions": 0,
            "textBoxes": 0,
            "headers": 0,
            "footers": 0,
            "footnotes": 0,
            "endnotes": 0,
            "comments": len(side["comments"]),
            "charts": chart_count,
            "images": images,
            "embeddedObjects": embedded,
            "hiddenSlides": hidden_slides,
            "groupedShapes": grouped_shapes,
            "notes": notes_count,
            "objects": chart_count + images + embedded,
            "unrendered": images + embedded,
        }
        return NativeText(
            "\n\n".join(texts), pages, flags["truncated"], {"office": office}
        )
    except Exception as cause:  # noqa: BLE001 - fall back to Docling
        print(f"python-pptx failed ({cause}); falling back to Docling", file=sys.stderr)
        return None


NATIVE_OFFICE = {
    ".docx": ("python-docx", extract_docx),
    ".xlsx": ("openpyxl", extract_xlsx),
    ".pptx": ("python-pptx", extract_pptx),
}


class DoclingEngine:
    name = "docling"

    def __init__(self) -> None:
        self._converter = self._build_converter()
        self.model = "docling"

    @staticmethod
    def _build_converter():
        from docling.document_converter import DocumentConverter, PdfFormatOption

        try:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions

            options = PdfPipelineOptions()
            # Enabling OCR explicitly keeps scanned PDFs and page images useful
            # even when an older Docling defaults to text-layer extraction only.
            options.do_ocr = True
            formats = {InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
            try:
                from docling.document_converter import ImageFormatOption

                formats[InputFormat.IMAGE] = ImageFormatOption(pipeline_options=options)
            except (ImportError, TypeError):
                # Older releases route images through the PDF pipeline defaults;
                # keep the PDF OCR options instead of discarding them.
                return DocumentConverter(format_options=formats)
            return DocumentConverter(format_options=formats)
        except Exception as cause:  # noqa: BLE001 - optional engine configuration
            print(f"docling OCR options unavailable ({cause}); using defaults", file=sys.stderr)
            return DocumentConverter()

    def extract(
        self,
        path: Path,
        mode: str = "auto",
        dpi: int | None = None,
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> Extraction:
        extension = path.suffix.lower()
        if extension in DOCLING_OFFICE_EXTENSIONS:
            validate_office_archive(path)
        # Office conversion reads the document itself; PDF/image conversion uses
        # the page-image pipeline (with OCR), so it must not claim a text layer.
        route = "native" if extension in DOCLING_OFFICE_EXTENSIONS else "visual"
        native = NATIVE_OFFICE.get(extension)
        if native is not None:
            label, extractor = native
            result = extractor(path)
            if result is not None:
                text, pages, native_truncated, native_meta = result
                if text.strip():
                    meta: dict[str, Any] = {"extractor": label}
                    if native_meta:
                        meta.update(native_meta)
                    document = dir_document("native", pages, meta)
                    return Extraction(
                        normalize_ocr_text(text),
                        None,
                        label,
                        self.model,
                        len(pages) or 1,
                        native_truncated
                        or (document.get("meta") or {}).get("truncated") is True,
                        "native",
                        document,
                    )

        max_pages = env_int("OCR_DOCLING_MAX_PAGES", DEFAULT_DOCLING_MAX_PAGES)
        try:
            result = self._converter.convert(
                str(path), page_range=(1, max_pages), max_num_pages=max_pages
            )
        except TypeError:  # older docling releases reject page bounds
            result = self._converter.convert(str(path))
        try:
            pages = len(result.document.pages)
        except Exception:  # noqa: BLE001 - page metadata is optional
            pages = 0
        markdown = result.document.export_to_markdown().strip()
        # The converter is page-bounded, so reaching the cap means the document
        # may have more pages than were exported.
        truncated = pages >= max_pages
        if len(markdown) > MAX_TEXT_CHARS:
            markdown = markdown[:MAX_TEXT_CHARS]
            truncated = True
        source = dir_source("docling", self.model, route)
        blocks = (
            [dir_block("p1-b0", "paragraph", markdown, None, None, source, 0)]
            if markdown
            else []
        )
        document = dir_document(
            route,
            [dir_page(1, None, None, blocks, source)],
            {"pages": pages, "truncated": truncated},
        )
        return Extraction(
            normalize_ocr_text(markdown),
            None,
            self.model,
            self.model,
            pages,
            truncated,
            route,
            document,
        )


ENGINES = {"paddle": PaddleEngine, "docling": DoclingEngine}
ENGINE_EXTENSIONS = {"paddle": PADDLE_EXTENSIONS, "docling": DOCLING_EXTENSIONS}


class Sidecar:
    def __init__(self, engine_name: str, upload_root: Path, token: str | None) -> None:
        self.engine_name = engine_name
        self.upload_root = upload_root.resolve()
        self.token = token
        self._engine: PaddleEngine | DoclingEngine | None = None
        self._lock = threading.Lock()

    def ensure_engine(self) -> PaddleEngine | DoclingEngine:
        with self._lock:
            if self._engine is None:
                started = time.perf_counter()
                self._engine = ENGINES[self.engine_name]()
                print(
                    f"{self.engine_name} engine ready in {time.perf_counter() - started:.1f}s",
                    file=sys.stderr,
                )
            return self._engine

    def resolve(self, raw: Any, name: Any) -> tuple[Path, str]:
        if not isinstance(raw, str) or not raw.strip():
            raise RequestError(400, "path is required")
        target = Path(raw)
        if not target.is_absolute():
            raise RequestError(400, "path must be absolute")
        resolved = target.resolve()
        if resolved != self.upload_root and self.upload_root not in resolved.parents:
            raise RequestError(403, "path escapes the upload root")
        if not resolved.is_file():
            raise RequestError(404, "file not found")
        # Stored uploads have no suffix, so classify by the original file name.
        extension = Path(name).suffix.lower() if isinstance(name, str) else ""
        extension = extension or resolved.suffix.lower()
        if extension not in ENGINE_EXTENSIONS[self.engine_name]:
            raise RequestError(
                400, f"{self.engine_name} does not handle {extension or 'this file'}"
            )
        return resolved, extension

    def extract(
        self,
        path: Path,
        extension: str,
        mode: str = "auto",
        dpi: int | None = None,
        lang: str | None = None,
        model: str | None = None,
        channel: str = "rgb",
    ) -> dict[str, Any]:
        if path.suffix.lower() == ".pdf" and pdf_needs_password(path):
            raise RequestError(422, "password_required: PDF is encrypted")
        engine = self.ensure_engine()
        started = time.perf_counter()
        with self._lock:
            if extension and path.suffix.lower() != extension:
                # PaddleX rejects inputs whose path lacks a known suffix.
                with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as handle:
                    temporary = Path(handle.name)
                try:
                    shutil.copyfile(path, temporary)
                    extraction = engine.extract(temporary, mode, dpi, lang, model, channel)
                finally:
                    temporary.unlink(missing_ok=True)
            else:
                extraction = engine.extract(path, mode, dpi, lang, model, channel)
        text = extraction.text
        confidence = extraction.confidence
        model_label_value = extraction.model
        base_model = extraction.base_model
        pages = extraction.pages
        truncated = extraction.truncated
        route = extraction.route
        document = extraction.document
        raw_text = raw_text_from_document(document)
        seconds = round(time.perf_counter() - started, 2)
        if not text.strip():
            raise RequestError(422, "no text recognized")
        if len(text) > MAX_TEXT_CHARS:
            text = text[:MAX_TEXT_CHARS]
            raw_text = raw_text[:MAX_TEXT_CHARS] if raw_text is not None else None
            truncated = True
        if isinstance(document, dict) and (document.get("meta") or {}).get("truncated") is True:
            truncated = True
        return {
            "text": text,
            "rawText": raw_text if raw_text is not None else text,
            "engine": model_label_value,
            "model": base_model,
            "confidence": round(confidence, 4) if confidence is not None else None,
            "pages": pages,
            "truncated": truncated,
            "route": route,
            "document": document,
            "preflight": preflight_image(path) if extension != ".pdf" else None,
            "seconds": seconds,
            "pipeline": PIPELINE_VERSION,
            "runtime": runtime_label(),
            "lang": lang,
            "channel": channel,
            "model_profile": model,
            "model_sha256": getattr(engine, "model_sha256", lambda _profile=None: None)(model),
            **runtime_details(),
        }

    def extract_tables(self, path: Path, extension: str) -> dict[str, Any]:
        """Optional Camelot table extraction for digital PDFs.

        Missing Camelot is not an error: the response reports `available:false`
        so the worker keeps the native table blocks.
        """
        if extension != ".pdf":
            raise RequestError(400, "table extraction supports PDFs only")
        try:
            import camelot  # type: ignore[import-not-found]
        except Exception as cause:  # noqa: BLE001 - optional dependency
            return {
                "available": False,
                "engine": "camelot",
                "tables": [],
                "reason": f"camelot unavailable: {cause}"[:300],
            }
        started = time.perf_counter()
        max_pages = env_int("OCR_CAMELOT_MAX_PAGES", DEFAULT_CAMELOT_MAX_PAGES)
        budget = env_int("OCR_CAMELOT_TIME_BUDGET_SECONDS", DEFAULT_CAMELOT_TIME_BUDGET_SECONDS)
        tables: list[dict[str, Any]] = []
        errors: list[str] = []
        truncated = False
        # Camelot parses the whole PDF in memory; serialise it with extraction
        # and bound both pages and wall-clock so one large file cannot exhaust
        # the sidecar.
        with self._lock:
            for flavor in ("lattice", "stream"):
                if time.perf_counter() - started >= budget:
                    truncated = True
                    break
                try:
                    found = camelot.read_pdf(
                        str(path), pages=f"1-{max_pages}", flavor=flavor
                    )
                except Exception as cause:  # noqa: BLE001 - try the next flavor
                    errors.append(f"{flavor}: {cause}"[:200])
                    continue
                for table in found:
                    if time.perf_counter() - started >= budget:
                        truncated = True
                        break
                    try:
                        rows = [
                            [
                                ""
                                if cell is None or (isinstance(cell, float) and math.isnan(cell))
                                else str(cell).strip()
                                for cell in row
                            ]
                            for row in table.df.values.tolist()
                        ]
                    except Exception:  # noqa: BLE001 - skip malformed tables
                        continue
                    if not any(any(cell for cell in row) for row in rows):
                        continue
                    if len(rows) > 500:
                        truncated = True
                    bbox: dict[str, float] | None = None
                    raw_bbox = getattr(table, "_bbox", None)
                    if isinstance(raw_bbox, (list, tuple)) and len(raw_bbox) == 4:
                        try:
                            x1, y1, x2, y2 = (float(value) for value in raw_bbox)
                            bbox = {
                                "x": min(x1, x2),
                                "y": min(y1, y2),
                                "width": abs(x2 - x1),
                                "height": abs(y2 - y1),
                            }
                        except (TypeError, ValueError):
                            bbox = None
                    tables.append(
                        {
                            "page": int(getattr(table, "page", 1)),
                            "rows": rows[:500],
                            "bbox": bbox,
                            "flavor": flavor,
                        }
                    )
                    if len(tables) >= 50:
                        truncated = True
                        break
                if tables or truncated:
                    break
        return {
            "available": True,
            "engine": "camelot",
            "tables": tables[:50],
            "truncated": truncated,
            "seconds": round(time.perf_counter() - started, 2),
            "errors": errors,
        }

    def extract_region(
        self,
        path: Path,
        extension: str,
        page_number: int,
        bbox: dict[str, Any],
        dpi: int | None,
        mode: str,
        padding: float,
        lang: str | None = None,
        model: str | None = None,
        bbox_dpi: int | None = None,
    ) -> dict[str, Any]:
        if self.engine_name != "paddle":
            raise RequestError(400, "region extraction is only available for PaddleOCR")
        if mode not in {"auto", "enhanced"}:
            raise RequestError(400, "region mode must be auto or enhanced")
        if extension != ".pdf":
            validate_image_budget(path)
        engine = self.ensure_engine()
        started = time.perf_counter()
        with self._lock:
            crop = crop_region_image(
                path, extension, page_number, bbox, dpi, padding, bbox_dpi
            )
            try:
                page_data = engine._page_from_image(  # type: ignore[attr-defined]
                    crop, mode, 1, lang, model
                )
            finally:
                crop.unlink(missing_ok=True)
        text = page_data.text.strip()
        if not text:
            raise RequestError(422, "no text recognized in region")
        return {
            "text": text,
            "confidence": (
                round(page_data.confidence, 4) if page_data.confidence is not None else None
            ),
            "engine": page_data.model,
            "model": getattr(
                engine, "_model_for", lambda _lang, _profile=None: getattr(engine, "model", None)
            )(lang, model),
            "seconds": round(time.perf_counter() - started, 2),
            "pipeline": PIPELINE_VERSION,
            "runtime": runtime_label(),
            "lang": lang,
            "model_profile": model,
            "model_sha256": getattr(engine, "model_sha256", lambda _profile=None: None)(model),
            **runtime_details(),
        }


class Handler(BaseHTTPRequestHandler):
    server_version = "PulseOcr/1.0"
    sidecar: Sidecar

    def log_message(self, format: str, *args: Any) -> None:  # noqa: A002 - stdlib signature
        if args and str(args[0]).startswith(("4", "5")):
            print(f"http {format % args}", file=sys.stderr)

    def send_json(self, status: int, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        try:
            self.wfile.write(body)
        except BrokenPipeError:
            pass

    def authorized(self) -> bool:
        expected = self.sidecar.token
        if not expected:
            return True
        provided = self.headers.get("X-OCR-Token", "")
        # compare_digest rejects non-ASCII str values; bytes always compare.
        return hmac.compare_digest(expected.encode("utf-8"), provided.encode("utf-8"))

    def do_GET(self) -> None:  # noqa: N802 - stdlib hook
        if self.path != "/health":
            self.send_json(404, {"error": "not found"})
            return
        if not self.authorized():
            self.send_json(401, {"error": "invalid token"})
            return
        payload = {
            "ok": True,
            "engine": self.sidecar.engine_name,
            "ready": self.sidecar._engine is not None,
            "pipeline": PIPELINE_VERSION,
            "runtime": runtime_label(),
            **runtime_details(),
        }
        if self.sidecar._engine is not None:
            payload["model"] = self.sidecar._engine.model
            payload["model_sha256"] = getattr(
                self.sidecar._engine, "model_sha256", lambda _profile=None: None
            )()
            profiles = getattr(self.sidecar._engine, "_profiles", None)
            if profiles:
                payload["model_profiles"] = sorted(profiles)
        self.send_json(200, payload)

    def do_POST(self) -> None:  # noqa: N802 - stdlib hook
        if self.path not in {"/extract", "/extract-region", "/tables"}:
            self.send_json(404, {"error": "not found"})
            return
        if not self.authorized():
            self.send_json(401, {"error": "invalid token"})
            return
        try:
            self.connection.settimeout(30)
        except OSError:
            pass
        if self.headers.get_content_type() != "application/json":
            self.send_json(415, {"error": "content-type must be application/json"})
            return
        length = self.headers.get("Content-Length")
        if not length or not length.isdigit():
            self.send_json(411, {"error": "content-length required"})
            return
        if int(length) > MAX_BODY_BYTES:
            self.send_json(413, {"error": "request body too large"})
            return
        try:
            payload = json.loads(self.rfile.read(int(length)) or b"{}")
        except (ValueError, json.JSONDecodeError):
            self.send_json(400, {"error": "invalid JSON body"})
            return
        if not isinstance(payload, dict):
            self.send_json(400, {"error": "invalid JSON body"})
            return

        try:
            path, extension = self.sidecar.resolve(payload.get("path"), payload.get("name"))
            if self.path == "/tables":
                self.send_json(200, self.sidecar.extract_tables(path, extension))
                return
            raw_mode = payload.get("mode", "auto")
            mode = raw_mode if isinstance(raw_mode, str) else "auto"
            if mode not in OCR_MODES:
                raise RequestError(
                    400, "mode must be auto, table, structure or enhanced"
                )
            dpi: int | None = None
            if payload.get("dpi") is not None:
                try:
                    dpi = int(payload["dpi"])
                except (TypeError, ValueError):
                    raise RequestError(400, "dpi must be an integer") from None
                if not 100 <= dpi <= 400:
                    raise RequestError(400, "dpi must be between 100 and 400")
            bbox_dpi: int | None = None
            if payload.get("bboxDpi") is not None:
                try:
                    bbox_dpi = int(payload["bboxDpi"])
                except (TypeError, ValueError):
                    raise RequestError(400, "bboxDpi must be an integer") from None
                if not 72 <= bbox_dpi <= 600:
                    raise RequestError(400, "bboxDpi must be between 72 and 600")
            if mode != "auto" and self.sidecar.engine_name != "paddle":
                raise RequestError(400, f"{mode} mode is only available for PaddleOCR")
            raw_channel = payload.get("channel", "rgb")
            channel = raw_channel if isinstance(raw_channel, str) else "rgb"
            if channel not in OCR_CHANNELS:
                raise RequestError(400, "channel must be rgb, gray or red")
            if channel != "rgb" and self.sidecar.engine_name != "paddle":
                raise RequestError(400, "channel is only available for PaddleOCR")
            raw_lang = payload.get("lang")
            lang: str | None = None
            if raw_lang is not None:
                if not isinstance(raw_lang, str) or raw_lang.strip().lower() not in OCR_LANGS:
                    raise RequestError(400, "lang must be one of ch, en, japan, korean, arabic, ru")
                lang = raw_lang.strip().lower()
                if self.sidecar.engine_name != "paddle":
                    raise RequestError(400, "lang is only available for PaddleOCR")
            if mode in {"table", "structure"} and extension not in PADDLE_EXTENSIONS:
                raise RequestError(
                    400, "table/structure mode supports images and PDFs"
                )
            raw_model = payload.get("model")
            model: str | None = None
            if raw_model is not None:
                if not isinstance(raw_model, str) or not raw_model.strip():
                    raise RequestError(400, "model must be a non-empty string")
                model = raw_model.strip()
                if self.sidecar.engine_name != "paddle":
                    raise RequestError(400, "model is only available for PaddleOCR")
                engine = self.sidecar.ensure_engine()
                profiles = getattr(engine, "_profiles", {})
                if model != "default" and model not in profiles:
                    raise RequestError(400, f"unknown model profile {model}")
            if self.path == "/extract-region":
                raw_page = payload.get("page", 1)
                try:
                    page_number = int(raw_page)
                except (TypeError, ValueError):
                    raise RequestError(400, "page must be an integer") from None
                if page_number < 1:
                    raise RequestError(400, "page must be >= 1")
                bbox = payload.get("bbox")
                if not isinstance(bbox, dict):
                    raise RequestError(400, "bbox is required")
                padding = payload.get("padding", 0.35)
                try:
                    padding_value = float(padding)
                except (TypeError, ValueError):
                    raise RequestError(400, "padding must be a number") from None
                if not 0 <= padding_value <= 1:
                    raise RequestError(400, "padding must be between 0 and 1")
                self.send_json(
                    200,
                    self.sidecar.extract_region(
                        path,
                        extension,
                        page_number,
                        bbox,
                        dpi,
                        mode,
                        padding_value,
                        lang,
                        model,
                        bbox_dpi,
                    ),
                )
                return
            self.send_json(
                200, self.sidecar.extract(path, extension, mode, dpi, lang, model, channel)
            )
        except RequestError as error:
            self.send_json(error.status, {"error": error.message})
        except Exception as cause:  # noqa: BLE001 - surface a JSON error to the worker
            traceback.print_exc()
            self.send_json(500, {"error": str(cause)[:500] or "extraction failed"})


def main() -> None:
    parser = argparse.ArgumentParser(description="Pulse local OCR sidecar")
    parser.add_argument("--engine", choices=sorted(ENGINES), required=True)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=None)
    parser.add_argument(
        "--upload-root",
        default=os.environ.get("OCR_UPLOAD_ROOT"),
        help="Only files under this directory are processed",
    )
    args = parser.parse_args()

    if not args.upload_root:
        parser.error("--upload-root or OCR_UPLOAD_ROOT is required")
    root = Path(args.upload_root)
    if not root.is_dir():
        parser.error(f"upload root does not exist: {root}")

    port = args.port if args.port is not None else (8765 if args.engine == "paddle" else 8766)
    token = os.environ.get("OCR_TOKEN") or None
    sidecar = Sidecar(args.engine, root, token)
    sidecar.ensure_engine()

    handler = type("BoundHandler", (Handler,), {"sidecar": sidecar})
    server = ThreadingHTTPServer((args.host, port), handler)
    server.daemon_threads = True
    print(
        f"ocr sidecar listening on http://{args.host}:{port} "
        f"engine={args.engine} root={root} auth={'on' if token else 'off'}",
        file=sys.stderr,
        flush=True,
    )
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
